from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "AGENTS.md").is_file())
TASK = Path(__file__).resolve().parents[1]
SOURCE = TASK / "report" / "REPORT_phase2_task01.md"
OUTPUT = TASK / "report" / "REPORT_phase2_task01.docx"


def set_font(run, name: str, size: float, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.font.bold = bold


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 110, bottom: int = 100, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def clean_inline(text: str) -> str:
    return text.replace("**", "").replace("`", "")


def remove_paragraph_borders(element) -> None:
    paragraph_properties = element.get_or_add_pPr()
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is not None:
        paragraph_properties.remove(border)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    data = [rows[0]] + rows[2:]
    columns = len(data[0])
    table = doc.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(data[0]):
        cell = table.rows[0].cells[index]
        cell.text = clean_inline(value)
        shade(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_font(run, "Arial", 9, True, "FFFFFF")
    for row_index, row in enumerate(data[1:]):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.text = clean_inline(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cell, "EAF2F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, columns - 1) else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_font(run, "Arial", 9)
    add_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("\n".join(lines))
    set_font(run, "Consolas", 9)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.75)

    styles = {"Normal": 11, "Title": 18, "Heading 1": 14, "Heading 2": 12}
    for name, size in styles.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.18
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for style_name in ("Title", "Heading 1", "Heading 2"):
        remove_paragraph_borders(doc.styles[style_name]._element)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    title_seen = False
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        index += 1
        if stripped.startswith("```"):
            if in_code:
                if "paragraph" in locals():
                    paragraph.paragraph_format.keep_with_next = True
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(raw)
            continue
        if not stripped:
            continue
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            current = stripped
            while True:
                rows.append([part.strip() for part in current.strip("|").split("|")])
                if index >= len(lines) or not lines[index].strip().startswith("|"):
                    break
                current = lines[index].strip()
                index += 1
            add_markdown_table(doc, rows)
            continue
        if stripped.startswith("# "):
            paragraph = doc.add_paragraph(clean_inline(stripped[2:]), style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            remove_paragraph_borders(paragraph._p)
            title_seen = True
            continue
        if stripped.startswith("## "):
            paragraph = doc.add_paragraph(clean_inline(stripped[3:]), style="Heading 1")
            paragraph.paragraph_format.keep_with_next = True
            continue
        if stripped.startswith("### "):
            paragraph = doc.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
            paragraph.paragraph_format.keep_with_next = True
            continue
        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(clean_inline(stripped[2:]))
        elif re.match(r"^\d+\.\s", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.add_run(clean_inline(re.sub(r"^\d+\.\s+", "", stripped)))
        else:
            paragraph = doc.add_paragraph(clean_inline(stripped))
        paragraph.paragraph_format.space_after = Pt(6)
        for run in paragraph.runs:
            set_font(run, "Arial", 11)

    if not title_seen:
        raise SystemExit("Report must contain a Markdown title")
    doc.core_properties.title = "Báo cáo Phase 2 Task 1 Soạn đặc tả cohort nháp"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
