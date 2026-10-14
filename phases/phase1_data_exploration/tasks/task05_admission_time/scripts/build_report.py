"""Build Task 5 Markdown and Word reports from a completed run."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "AGENTS.md").is_file())


def fmt(value: int | float, digits: int = 0) -> str:
    return f"{value:,.{digits}f}"


def pct(value: int | float, total: int | float) -> str:
    return f"{100 * value / total:.3f}%" if total else "—"


def markdown_table(headers, rows):
    return ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"] + ["| " + " | ".join(map(str, row)) + " |" for row in rows]


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); shd = tc_pr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None: tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None: node = OxmlElement(f"w:{m}"); tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = str(header); cell.width = Cm(widths[index]); shade(cell, "1F4E78"); set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(9); run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row)):
            cell.text = str(value); cell.width = Cm(widths[index]); set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2: shade(cell, "EAF2F8")
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(9)
    tbl_pr = table._tbl.tblPr; borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "D9D9D9"); borders.append(el)
    tbl_pr.append(borders)
    return table


def build_docx(path: Path, sections: list[tuple]) -> None:
    template = ROOT / "tmp" / "phase1_reference.docx"
    doc = Document(template) if template.exists() else Document()
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"): body.remove(child)
    for section in doc.sections:
        section.top_margin = Cm(2); section.bottom_margin = Cm(2); section.left_margin = Cm(2.4); section.right_margin = Cm(2)
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            for paragraph in part.paragraphs: paragraph.clear()
    for name, size in (("Normal", 11), ("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[name]; style.font.name = "Arial"; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7); style.paragraph_format.line_spacing = 1.2
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for entry in sections:
        kind = entry[0]
        if kind == "title": doc.add_paragraph(entry[1], "Title")
        elif kind == "subtitle":
            p = doc.add_paragraph(entry[1]); p.paragraph_format.space_after = Pt(12)
        elif kind == "h1": doc.add_paragraph(entry[1], "Heading 1")
        elif kind == "h2": doc.add_paragraph(entry[1], "Heading 2")
        elif kind == "p": doc.add_paragraph(entry[1])
        elif kind == "bullet":
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.add_run("• " + entry[1])
        elif kind == "table": add_table(doc, entry[1], entry[2], entry[3])
    doc.core_properties.title = "Báo cáo Task 5 khảo sát admission và thời gian"; doc.core_properties.author = ""
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("run_dir", type=Path); args = parser.parse_args(); run = args.run_dir.resolve()
    data = json.loads((run / "results.json").read_text(encoding="utf-8")); c = data["counts"]
    windows = pd.read_csv(run / "window_coverage.csv"); rel = pd.read_csv(run / "lab_time_distribution.csv"); lag = pd.read_csv(run / "storetime_lag_distribution.csv")
    admission_rows = [
        ["Tổng admission", fmt(data["admissions"]), "100.000%"],
        ["Interval thời gian hợp lệ", fmt(data["valid_interval_admissions"]), pct(data["valid_interval_admissions"], data["admissions"])],
        ["dischtime trước admittime", fmt(data["negative_duration_admissions"]), pct(data["negative_duration_admissions"], data["admissions"])],
        ["Thời lượng bằng 0", fmt(data["zero_duration_admissions"]), pct(data["zero_duration_admissions"], data["admissions"])],
    ]
    position_rows = [
        ["Dòng ứng viên", fmt(c["candidate_rows"]), "100.000%"],
        ["Trước admittime", fmt(c["before_admission_rows"]), pct(c["before_admission_rows"], c["candidate_rows"])],
        ["Trong admission", fmt(c["during_admission_rows"]), pct(c["during_admission_rows"], c["candidate_rows"])],
        ["Sau dischtime", fmt(c["after_discharge_rows"]), pct(c["after_discharge_rows"], c["candidate_rows"])],
        ["Admission có interval âm hoặc thiếu", fmt(c["invalid_interval_rows"]), pct(c["invalid_interval_rows"], c["candidate_rows"])],
    ]
    window_rows = [[f"{int(r.window_hours)} giờ", fmt(r.lab_rows), fmt(r.admissions_with_lab), f"{100*r.share_all_admissions:.3f}%", fmt(r.admissions_discharged_before_window)] for r in windows.itertuples()]
    store_rows = [[r.storetime_lag_bin, fmt(r.rows), f"{100*r.share_valid_storetime:.3f}%"] for r in lag.itertuples() if r.rows]
    q = data["first_in_admission_lab_hour"]
    md = [
        "# Báo cáo Task 5 khảo sát admission và thời gian", "", "## 1 Mục đích và phạm vi", "",
        "Task 5 đo quan hệ giữa thời gian xét nghiệm và thời gian admission trên toàn bộ dòng đạt định nghĩa nền tảng sử dụng của Task 3. Các dòng thiếu `hadm_id` không được khôi phục. Dữ liệu HOSP gốc chỉ được đọc.", "",
        "Các cửa sổ 6, 12 và 24 giờ là so sánh mô tả. Cửa sổ dùng quy ước `0 <= charttime - admittime < số giờ` và chỉ tính dòng nằm trong admission có interval không âm. Task chưa chọn cửa sổ chính thức.", "",
        "## 2 Chất lượng thời gian admission", "", *markdown_table(["Chỉ số", "Số admission", "Tỷ lệ"], admission_rows), "",
        f"Thời lượng admission hợp lệ có trung vị {fmt(data['admission_duration_hours']['median'], 2)} giờ, phân vị 25 là {fmt(data['admission_duration_hours']['p25'], 2)} giờ và phân vị 75 là {fmt(data['admission_duration_hours']['p75'], 2)} giờ.", "",
        "## 3 Vị trí thời gian của xét nghiệm", "", *markdown_table(["Nhóm", "Số dòng", "Tỷ lệ"], position_rows), "",
        f"Có {fmt(data['candidate_admissions'])} admission có ít nhất một dòng ứng viên; {fmt(data['candidate_admissions_with_during_lab'])} admission có xét nghiệm nằm trong interval admission hợp lệ. Có {fmt(data['candidate_admissions_only_outside_valid_interval'])} admission chỉ có dòng ứng viên ngoài interval hợp lệ hoặc thuộc admission có thời lượng bất thường.", "",
        "## 4 Độ phủ cửa sổ mô tả", "", *markdown_table(["Cửa sổ", "Dòng lab", "Admission có lab", "Tỷ lệ toàn admission", "Admission kết thúc trước cửa sổ"], window_rows), "",
        "Các cửa sổ dài hơn giữ thêm xét nghiệm và admission, đồng thời tăng khả năng chứa thông tin xuất hiện muộn hơn trong quá trình điều trị. Số liệu này chưa đủ để tự động chọn một cửa sổ.", "",
        "## 5 Thời điểm xét nghiệm đầu tiên", "",
        f"Trong các admission có xét nghiệm nằm trong admission, thời điểm xét nghiệm đầu tiên tính từ admittime có trung vị {fmt(q['median'], 2)} giờ; phân vị 25 là {fmt(q['p25'], 2)} giờ; phân vị 75 là {fmt(q['p75'], 2)} giờ; phân vị 95 là {fmt(q['p95'], 2)} giờ.", "",
        "## 6 Thời gian lưu kết quả", "", *markdown_table(["Khoảng storetime trừ charttime", "Số dòng", "Tỷ lệ trên storetime hợp lệ"], store_rows), "",
        f"Có {fmt(c.get('candidate_missing_storetime', 0))} dòng ứng viên thiếu storetime và {fmt(c.get('candidate_storetime_before_charttime', 0))} dòng có storetime trước charttime. Storetime chưa được chọn làm mốc feature.", "",
        "## 7 Kết luận Task 5", "",
        "Task 5 đã xác định chính xác số dòng và admission còn lại trong từng cửa sổ mô tả. Task chưa làm sạch dữ liệu, chưa chọn cohort và chưa quyết định 6, 12 hay 24 giờ. Các interval admission âm và các dòng ngoài thời gian nằm viện được chuyển sang Task 6 để quản lý như vấn đề chất lượng.", "",
        "## 8 File đầu ra", "", "- `admission_time_profile.csv`: hồ sơ thời gian và số dòng lab trên từng admission.", "- `window_coverage.csv`: độ phủ 6, 12 và 24 giờ.", "- `lab_time_distribution.csv`: phân bố dòng lab theo khoảng từ admittime.", "- `storetime_lag_distribution.csv`: phân bố độ trễ lưu kết quả.", "- `results.json`: metadata và số liệu máy đọc.", "- `task05_time_review.xlsx`: workbook dành cho người đọc.",
    ]
    (run / "REPORT_task05.md").write_text("\n".join(md), encoding="utf-8")
    sections = [
        ("title", "BÁO CÁO TASK 5 KHẢO SÁT THỜI GIAN"), ("subtitle", "Giai đoạn 1 khảo sát admission và dữ liệu xét nghiệm MIMIC IV"),
        ("p", "Báo cáo đo vị trí thời gian của xét nghiệm so với admission và độ phủ mô tả trong 6, 12 và 24 giờ đầu. Các dòng thiếu hadm_id không được khôi phục và dữ liệu nguồn không bị sửa."),
        ("h1", "1 Mục đích và phạm vi"), ("p", "Task 5 sử dụng toàn bộ dòng đáp ứng định nghĩa nền tảng sử dụng của Task 3. Ba cửa sổ chỉ dùng để so sánh độ phủ; báo cáo chưa chọn cửa sổ chính thức hoặc cohort cuối."),
        ("h1", "2 Chất lượng thời gian admission"), ("table", ["Chỉ số", "Số admission", "Tỷ lệ"], admission_rows, [8.5, 3.5, 3.5]),
        ("p", f"Thời lượng hợp lệ có trung vị {fmt(data['admission_duration_hours']['median'], 2)} giờ. Admission có dischtime trước admittime không được dùng để phân loại lab trước, trong hoặc sau admission."),
        ("h1", "3 Vị trí thời gian của xét nghiệm"), ("table", ["Nhóm", "Số dòng", "Tỷ lệ"], position_rows, [8.5, 3.5, 3.5]),
        ("p", f"Có {fmt(data['candidate_admissions_with_during_lab'])} admission có ít nhất một xét nghiệm nằm trong interval admission hợp lệ."),
        ("h1", "4 Độ phủ cửa sổ mô tả"), ("table", ["Cửa sổ", "Dòng lab", "Admission", "Tỷ lệ", "Kết thúc sớm"], window_rows, [2.2, 3.1, 3.1, 2.7, 3.5]),
        ("p", "Quy ước cửa sổ là nửa mở từ admittime và chỉ tính dòng nằm trong admission. Cửa sổ dài hơn tăng độ phủ nhưng cũng nhận thêm thông tin xuất hiện muộn hơn trong điều trị."),
        ("h1", "5 Thời điểm xét nghiệm đầu tiên"), ("p", f"Thời điểm xét nghiệm đầu tiên có trung vị {fmt(q['median'], 2)} giờ từ admittime; phân vị 25 là {fmt(q['p25'], 2)} giờ và phân vị 75 là {fmt(q['p75'], 2)} giờ."),
        ("h1", "6 Thời gian lưu kết quả"), ("table", ["Khoảng thời gian", "Số dòng", "Tỷ lệ"], store_rows, [8.5, 3.5, 3.5]),
        ("p", f"Có {fmt(c.get('candidate_missing_storetime', 0))} dòng ứng viên thiếu storetime và {fmt(c.get('candidate_storetime_before_charttime', 0))} dòng có storetime trước charttime. Charttime vẫn là mốc chính đang khảo sát."),
        ("h1", "7 Kết luận"), ("p", "Task 5 cung cấp số liệu để thảo luận cửa sổ sớm ở Phase 2. Task chưa chọn 6, 12 hay 24 giờ và chưa tạo dataset huấn luyện."),
        ("h1", "8 File đầu ra"), ("p", "Kết quả gồm admission_time_profile.csv, window_coverage.csv, lab_time_distribution.csv, storetime_lag_distribution.csv, results.json và task05_time_review.xlsx."),
    ]
    build_docx(run / "REPORT_task05.docx", sections); print(run / "REPORT_task05.docx")


if __name__ == "__main__": main()
