from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"
OUTPUT_DIR = ROOT / "outputs"
MD_PATH = REPORT_DIR / "REPORT_task07.md"
DOCX_PATH = REPORT_DIR / "REPORT_task07.docx"
DECISION_PATH = OUTPUT_DIR / "decision_register_task07.csv"
PLAN_PATH = OUTPUT_DIR / "phase2_proposed_plan_task07.csv"


TASKS = [
    ("1", "Inventory HOSP", "Hoàn tất", "Kiểm kê 22 bảng, schema và quy mô; bổ sung từ điển tiếng Việt cho 44 cột của 5 bảng chính."),
    ("2", "Khóa và liên kết", "Hoàn tất", "Kiểm tra toàn bộ khóa, orphan, sai bệnh nhân và bốn nhóm overlap admission."),
    ("3", "Khảo sát labevents", "Hoàn tất", "Đánh giá 158.374.764 dòng về khả năng liên kết, kết quả, đơn vị và độ phủ itemid."),
    ("4", "Khảo sát diagnoses_icd", "Hoàn tất", "Đánh giá 6.364.488 dòng ICD-9/10, dictionary, trùng mã và phân bố diagnosis/admission."),
    ("5", "Admission và thời gian", "Hoàn tất", "Đo quan hệ charttime/admittime/dischtime và độ phủ mô tả trong 6/12/24 giờ."),
    ("6", "Chất lượng và leakage", "Hoàn tất", "Lập danh mục 11 vấn đề chất lượng, 6 nguy cơ leakage và mức sẵn sàng của cohort."),
    ("7", "Tổng hợp Phase 1", "Hoàn tất", "Tổng hợp kết quả, quyết định, vấn đề còn mở và đề xuất kế hoạch Phase 2."),
]

CORE_TABLES = [
    ("admissions", "546.028", "546.028 admission", "Xác định admission và mốc nhập/xuất viện"),
    ("patients", "364.627", "364.627 bệnh nhân", "Đối chiếu bệnh nhân; chia tập theo subject_id"),
    ("labevents", "158.374.764", "Kết quả xét nghiệm", "Nguồn feature X"),
    ("d_labitems", "1.650", "Từ điển xét nghiệm", "Tên, bệnh phẩm, category của itemid"),
    ("diagnoses_icd", "6.364.488", "Mã ICD theo admission", "Nguồn nhãn Y trước ánh xạ PheCode"),
    ("d_icd_diagnoses", "112.107 khóa", "Từ điển ICD", "Mô tả mã theo icd_code + icd_version"),
]

LINKAGE = [
    ("Có cả lab và diagnosis", "447.414", "81,940%", "Điểm xuất phát liên kết; chưa phải cohort cuối"),
    ("Chỉ có lab", "274", "0,050%", "Không có nhãn diagnosis cho thiết kế hiện tại"),
    ("Chỉ có diagnosis", "98.083", "17,963%", "Không có feature lab liên kết trực tiếp"),
    ("Không có cả hai", "257", "0,047%", "Không dùng cho bài toán X + Y hiện tại"),
]

LAB_FUNNEL = [
    ("Tổng labevents", "158.374.764", "100,000%"),
    ("Có hadm_id", "84.605.867", "53,421%"),
    ("Khớp admission và subject_id", "84.605.454", "53,421%"),
    ("Có kết quả số", "136.884.423", "86,431%"),
    ("Chỉ có kết quả chữ", "5.427.558", "3,427%"),
    ("Không có kết quả số hoặc chữ", "16.062.783", "10,142%"),
    ("Đủ nền tảng sử dụng của Task 3", "79.439.965", "50,159%"),
]

WINDOWS = [
    ("6 giờ", "4.701.358", "189.225", "189.102", "34,632%"),
    ("12 giờ", "9.786.000", "308.886", "308.719", "56,539%"),
    ("24 giờ", "16.805.692", "400.315", "400.092", "73,273%"),
]

QUALITY = [
    ("labevents", "Thiếu hadm_id", "73.768.897 dòng; 46,579%", "Loại khỏi dữ liệu xử lý admission-level; không khôi phục khóa"),
    ("labevents", "subject_id sai so với admission", "413 dòng", "Loại khỏi tập ứng viên; lưu audit count"),
    ("labevents", "Không có kết quả số hoặc chữ", "16.062.783 dòng; 10,142%", "Không dùng làm giá trị feature"),
    ("labevents", "Thiếu storetime", "224.118 dòng ứng viên", "Không loại chỉ vì thiếu storetime; tiếp tục kiểm soát tính sẵn có"),
    ("labevents", "storetime trước charttime", "762 dòng", "Kiểm tra trước khi dùng storetime"),
    ("lab item", "Nhiều chuỗi đơn vị", "10/976 itemid", "Xem xét từng itemid trước chuẩn hóa"),
    ("admissions", "dischtime trước admittime", "175 admission", "Không dùng interval âm cho cohort/window"),
    ("lab/admission", "Lab trước admittime", "2.816.594 dòng", "Không đưa vào cửa sổ bắt đầu tại admittime nếu chưa có quy tắc khác"),
    ("lab/admission", "Lab sau dischtime", "53.761 dòng", "Không đưa vào feature của admission"),
    ("diagnoses_icd", "Trùng mã trong admission", "388 dòng dư", "Khử trùng khi tạo nhãn admission-phenotype"),
    ("diagnoses_icd", "Nhiều seq_num = 1", "10 admission", "Không giả định primary diagnosis luôn duy nhất"),
]

DECISIONS = [
    ("D01", "Đã chốt", "Đơn vị phân tích", "Một hàng cuối cùng đại diện một hadm_id."),
    ("D02", "Đã chốt", "Chia tập", "Chia train/validation/test theo subject_id."),
    ("D03", "Đã chốt", "Lab thiếu hadm_id", "Không khôi phục hoặc suy diễn; loại khỏi dataset xử lý admission-level và ghi audit."),
    ("D04", "Đã chốt", "Cách ghép X và Y", "Tổng hợp nhánh lab và nhánh nhãn riêng ở cấp hadm_id rồi mới ghép."),
    ("D05", "Đã chốt", "Diễn giải diagnosis", "Không coi diagnoses_icd là timestamp chẩn đoán và không tuyên bố dự đoán trước bác sĩ."),
    ("D06", "Chờ quyết định", "Early-admission window", "Chọn 6, 12 hoặc 24 giờ sau khi cân bằng độ phủ và nguy cơ thông tin muộn."),
    ("D07", "Chờ quyết định", "Cohort", "Chốt inclusion/exclusion, cách xử lý admission interval âm và admission kết thúc sớm."),
    ("D08", "Chờ quyết định", "Đơn vị xét nghiệm", "Quy tắc cho 10 itemid có nhiều chuỗi đơn vị."),
    ("D09", "Chờ quyết định", "Giá trị xét nghiệm", "Quy tắc dùng kết quả chữ, missing và outlier."),
    ("D10", "Chờ khảo sát Phase 4", "Nhãn phenotype", "Map ICD→PheCode, đo coverage và frequency rồi mới chọn threshold/label set."),
]

LEAKAGE = [
    ("Xét nghiệm muộn", "Chọn cửa sổ sớm đã được duyệt; chỉ tạo feature trong cửa sổ đó."),
    ("Cùng bệnh nhân ở nhiều tập", "Group split theo subject_id."),
    ("Tiền xử lý học từ toàn bộ dữ liệu", "Fit imputation, scaling và lựa chọn học từ dữ liệu chỉ trên train."),
    ("Chọn nhãn dựa trên test", "Chốt protocol lựa chọn nhãn trước đánh giá cuối."),
    ("Diễn giải sai thời điểm diagnosis", "Chỉ nói phenotype được ghi nhận cho admission."),
    ("Nhân dòng khi join", "Tổng hợp X và Y riêng theo hadm_id, sau đó mới nối."),
]

PHASE2_PLAN = [
    ("2.1", "Soạn cohort specification", "Tập hợp tiêu chí bắt buộc và lựa chọn còn mở; định nghĩa mốc thời gian và biên cửa sổ.", "Bản đặc tả cohort để duyệt", "Chờ duyệt"),
    ("2.2", "So sánh kịch bản cohort", "Lập bảng 6/12/24 giờ theo admission, subject, số lab, thời lượng và nhóm bị loại.", "Bảng so sánh có số liệu", "Đề xuất"),
    ("2.3", "Chốt inclusion/exclusion", "Quyết định interval âm, admission kết thúc sớm, lab ngoài admission và dòng sai subject_id.", "Decision log và quy tắc lọc", "Chờ quyết định"),
    ("2.4", "Tạo cohort xử lý mới", "Áp dụng quy tắc đã duyệt, không ghi đè nguồn và ghi funnel trước/sau từng bước.", "Danh sách hadm_id/subject_id và audit counts", "Sau 2.1–2.3"),
    ("2.5", "Kiểm tra và đóng băng cohort", "Kiểm tra khóa duy nhất, liên kết, thời gian, overlap và khả năng tái lập.", "Cohort manifest/version", "Sau 2.4"),
]

DELIVERABLES = [
    ("Báo cáo Task 1–6", "06 bộ Markdown/Word cùng các bảng CSV/XLSX theo từng task", "Bằng chứng chi tiết cho từng bước khảo sát"),
    ("Từ điển dữ liệu", "44 cột của 5 bảng chính, có bản Word, Excel và CSV", "Giải thích schema và vai trò của từng trường"),
    ("Báo cáo tổng kết", "REPORT_task07.md và REPORT_task07.docx", "Báo cáo định kỳ cho giảng viên"),
    ("Nhật ký quyết định", "decision_register_task07.csv", "Tách nội dung đã chốt và nội dung chờ duyệt"),
    ("Kế hoạch Phase 2", "phase2_proposed_plan_task07.csv", "Khung thảo luận cho bước định nghĩa cohort"),
    ("Mã tái lập", "Script của từng task và manifest lần chạy", "Cho phép kiểm tra lại số liệu mà không sửa dữ liệu gốc"),
]


def write_csv(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def write_markdown():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Báo cáo định kỳ — Tổng kết Giai đoạn 1: Khảo sát dữ liệu MIMIC-IV HOSP", "",
        "**Ngày báo cáo:** 04/10/2026  ", "**Phạm vi:** Task 1 đến Task 7  ",
        "**Dữ liệu:** `D:\\UIT_thu\\NCKH\\data\\hosp` (chỉ đọc)", "",
        "## 1. Mục tiêu nghiên cứu", "",
        "Nghiên cứu xây dựng mô hình Machine Learning sử dụng dữ liệu xét nghiệm giai đoạn sớm của một hospital admission để dự đoán đồng thời các phenotype bệnh được ghi nhận cho admission đó. Đơn vị phân tích là `hadm_id`; bài toán là multi-label classification; chia tập theo `subject_id`.", "",
        "## 2. Tóm tắt điều hành", "",
        "- Hoàn tất 7/7 task của Phase 1 và không sửa dữ liệu HOSP gốc.",
        "- Đã khảo sát 22 bảng; phân tích sâu các bảng admission, lab và diagnosis.",
        "- Có 446.891 admission có cả lab ứng viên và diagnosis; 432.410 admission có lab nằm trong admission hợp lệ và có diagnosis.",
        "- Các kịch bản 6/12/24 giờ có lần lượt 189.102/308.719/400.092 admission có cả lab và diagnosis.",
        "- Đã chốt loại lab thiếu `hadm_id` khỏi dataset xử lý admission-level, tổng hợp X/Y riêng và chia tập theo bệnh nhân.",
        "- Chưa chọn cửa sổ chính thức, cohort cuối, quy tắc đơn vị/giá trị hoặc label set PheCode.", "",
        "## 3. Tiến độ Task 1–7", "", "| Task | Nội dung | Trạng thái | Kết quả chính |", "|---:|---|---|---|",
    ]
    lines += ["| " + " | ".join(x) + " |" for x in TASKS]
    lines += ["", "## 4. Quy mô và cấu trúc dữ liệu", "", "| Bảng | Quy mô | Đơn vị dòng | Vai trò |", "|---|---:|---|---|"]
    lines += ["| " + " | ".join(x) + " |" for x in CORE_TABLES]
    lines += ["", "## 5. Khả năng liên kết ở cấp admission", "", "| Nhóm admission | Số admission | Tỷ lệ | Ý nghĩa |", "|---|---:|---:|---|"]
    lines += ["| " + " | ".join(x) + " |" for x in LINKAGE]
    lines += ["", "Có 413 dòng lab sai `subject_id` so với chủ sở hữu `hadm_id`. Nếu nối trực tiếp nhiều lab với nhiều diagnosis, phép nối dự kiến tạo khoảng 1,685 tỷ dòng; vì vậy hai nhánh phải được tổng hợp riêng.", "", "## 6. Kết quả khảo sát labevents", "", "| Chỉ số | Số dòng | Tỷ lệ |", "|---|---:|---:|"]
    lines += ["| " + " | ".join(x) + " |" for x in LAB_FUNNEL]
    lines += ["", "Có 447.165 admission và 817 itemid có ít nhất một dòng đạt định nghĩa nền tảng sử dụng. Có 10 itemid xuất hiện với nhiều chuỗi đơn vị. Trung vị số dòng xét nghiệm trên một admission là 85; thống kê này chưa giới hạn cửa sổ sớm.", "", "## 7. Kết quả khảo sát diagnosis", "", "- Toàn bộ 6.364.488 dòng đủ khóa, khớp admission, bệnh nhân và dictionary.", "- Có 545.497 admission, 2.908.741 dòng ICD-9 và 3.455.747 dòng ICD-10.", "- Có 9.143 mã ICD-9, 19.440 mã ICD-10 và 28.583 cặp `icd_code + icd_version`.", "- Trung vị 10 diagnosis/admission; lớn nhất 57 dòng diagnosis.", "- Có 388 dòng dư trùng mã trong admission và 10 admission có nhiều hơn một `seq_num = 1`.", "- Dữ liệu đủ nền tảng để ánh xạ PheCode, nhưng chưa thể coi toàn bộ ICD là label set cuối.", "", "## 8. Thời gian và ba kịch bản cửa sổ", "", "Quy ước khảo sát là `0 <= charttime - admittime < số giờ`, chỉ trên admission có interval không âm.", "", "| Cửa sổ | Dòng lab | Admission có lab | Admission có lab + diagnosis | Tỷ lệ toàn admission |", "|---|---:|---:|---:|---:|"]
    lines += ["| " + " | ".join(x) + " |" for x in WINDOWS]
    lines += ["", "Cửa sổ dài hơn tăng độ phủ nhưng cũng tăng khả năng sử dụng thông tin xuất hiện muộn. Phase 1 chưa lựa chọn 6, 12 hay 24 giờ.", "", "## 9. Danh mục vấn đề chất lượng", "", "| Nguồn | Vấn đề | Mức ảnh hưởng | Hướng xử lý/điểm cần chốt |", "|---|---|---:|---|"]
    lines += ["| " + " | ".join(x) + " |" for x in QUALITY]
    lines += ["", "## 10. Kiểm soát leakage", "", "| Nguy cơ | Kiểm soát |", "|---|---|"]
    lines += ["| " + " | ".join(x) + " |" for x in LEAKAGE]
    lines += ["", "## 11. Quyết định và vấn đề còn mở", "", "| Mã | Trạng thái | Chủ đề | Nội dung |", "|---|---|---|---|"]
    lines += ["| " + " | ".join(x) + " |" for x in DECISIONS]
    lines += ["", "## 12. Đánh giá kết thúc Phase 1", "", "Phase 1 đã hoàn thành mục tiêu khảo sát cấu trúc, độ đầy đủ, liên kết, thời gian và rủi ro dữ liệu. Dữ liệu chưa phải dataset sạch để huấn luyện vì cohort và các quy tắc biến đổi chưa được chốt. Có đủ bằng chứng để bắt đầu thiết kế Phase 2 ở mức đặc tả và so sánh kịch bản.", "", "## 13. Đề xuất kế hoạch Phase 2 — Cohort Definition", "", "Đây là kế hoạch đề xuất để thảo luận; chưa phải quyết định nghiên cứu đã chốt.", "", "| Task | Nội dung | Công việc | Đầu ra | Trạng thái |", "|---|---|---|---|---|"]
    lines += ["| " + " | ".join(x) + " |" for x in PHASE2_PLAN]
    lines += ["", "### Các quyết định cần thống nhất trước khi tạo cohort", "", "1. Mục tiêu cân bằng giữa độ phủ và tính sớm khi chọn 6/12/24 giờ.", "2. Xử lý admission có `dischtime < admittime` và admission kết thúc trước độ dài cửa sổ.", "3. Có yêu cầu tối thiểu về loại xét nghiệm hoặc số lượng kết quả trong cửa sổ hay không.", "4. Cách kiểm soát `storetime` khi kết quả được lưu sau mốc dự đoán.", "5. Có giới hạn quần thể bổ sung theo tuổi, loại admission hoặc đặc điểm khác hay không; chỉ bổ sung nếu phục vụ câu hỏi nghiên cứu.", "", "## 14. Kết luận báo cáo định kỳ", "", "Giai đoạn 1 đã tạo nền tảng có thể tái lập cho việc định nghĩa cohort: dữ liệu nguồn được giữ nguyên, các bảng và khóa chính đã được xác minh, độ phủ của các kịch bản thời gian đã được đo, và các nguy cơ leakage đã được ghi thành quy tắc kiểm soát. Công việc tiếp theo nên bắt đầu bằng việc duyệt đặc tả cohort, sau đó mới tạo tập dữ liệu xử lý mới cho Phase 2.", "", "## Phụ lục A. Sản phẩm đã hoàn thành", "", "| Nhóm sản phẩm | Thành phần | Mục đích |", "|---|---|---|"]
    lines += ["| " + " | ".join(x) + " |" for x in DELIVERABLES]
    lines += ["", "## Phụ lục B. Nội dung đề nghị trao đổi với giảng viên", "", "- Xác nhận phạm vi bài toán multi-label ở cấp admission và cách diễn giải đầu ra.", "- Thảo luận tiêu chí chọn cửa sổ 6/12/24 giờ dựa trên độ phủ và nguy cơ thông tin muộn.", "- Thống nhất inclusion/exclusion criteria trước khi tạo cohort xử lý.", "- Xác nhận lộ trình: Phase 2 cohort, Phase 3 lab features, Phase 4 PheCode labels, sau đó mới ghép dataset và modeling.", "- Ghi nhận rằng Phase 1 mới khảo sát và kiểm soát rủi ro; chưa có kết quả mô hình."]
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_cell(cell, value, bold=False, white=False, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value))
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.bold = bold
    if white:
        r.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, white=True, size=font_size)
        set_cell_shading(table.rows[0].cells[i], "1F4E79")
        if widths:
            table.rows[0].cells[i].width = Cm(widths[i])
    for ridx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell(cells[i], value, size=font_size)
            if ridx % 2:
                set_cell_shading(cells[i], "EAF2F8")
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def style_document(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 12)]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(31, 78, 121)
    styles["Title"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def write_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("BÁO CÁO ĐỊNH KỲ\nTỔNG KẾT GIAI ĐOẠN 1")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Khảo sát dữ liệu MIMIC-IV v3.1 HOSP")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(89, 89, 89)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Task 1–7 | Ngày báo cáo: 04/10/2026").italic = True
    p = doc.add_paragraph()
    p.add_run("Dữ liệu nguồn: ").bold = True
    p.add_run(r"D:\UIT_thu\NCKH\data\hosp (chỉ đọc; không ghi đè).")

    doc.add_heading("1. Mục tiêu nghiên cứu", level=1)
    doc.add_paragraph("Nghiên cứu hướng đến xây dựng mô hình Machine Learning sử dụng dữ liệu xét nghiệm cận lâm sàng ở giai đoạn sớm của một hospital admission để dự đoán đồng thời các phenotype bệnh được ghi nhận cho admission đó. Đơn vị phân tích là hadm_id, bài toán là multi-label classification và việc chia tập phải thực hiện theo subject_id.")
    doc.add_heading("2. Tóm tắt điều hành", level=1)
    add_bullets(doc, ["Hoàn tất 7/7 task của Phase 1; dữ liệu HOSP gốc không bị sửa.", "Đã kiểm kê 22 bảng và phân tích sâu dữ liệu admission, xét nghiệm và chẩn đoán.", "Có 446.891 admission có cả lab ứng viên và diagnosis; 432.410 admission có lab nằm trong admission hợp lệ và có diagnosis.", "Ba kịch bản 6/12/24 giờ có lần lượt 189.102/308.719/400.092 admission có cả lab và diagnosis.", "Đã chốt loại lab thiếu hadm_id khỏi dataset xử lý, tổng hợp X/Y riêng theo hadm_id và chia tập theo subject_id.", "Chưa chọn cửa sổ chính thức, cohort cuối, quy tắc đơn vị/giá trị hoặc label set PheCode."])
    doc.add_heading("3. Tiến độ Task 1–7", level=1)
    add_table(doc, ["Task", "Nội dung", "Trạng thái", "Kết quả chính"], TASKS, [1.2, 3.5, 2.2, 10.0], 7.5)
    doc.add_heading("4. Quy mô và cấu trúc dữ liệu", level=1)
    add_table(doc, ["Bảng", "Quy mô", "Đơn vị dòng", "Vai trò"], CORE_TABLES, [3.2, 2.8, 4.2, 6.8], 8)
    doc.add_paragraph("Các bảng medication, order, procedure và dữ liệu ECG không phải nguồn input chính của hướng nghiên cứu hiện tại.")
    doc.add_heading("5. Khả năng liên kết ở cấp admission", level=1)
    add_table(doc, ["Nhóm admission", "Số admission", "Tỷ lệ", "Ý nghĩa"], LINKAGE, [4.4, 2.7, 2.0, 7.9], 8)
    doc.add_paragraph("Có 413 dòng lab sai subject_id so với chủ sở hữu hadm_id. Nếu nối trực tiếp nhiều lab với nhiều diagnosis, phép nối dự kiến tạo khoảng 1,685 tỷ dòng. Pipeline phải tổng hợp X và Y riêng ở cấp hadm_id rồi mới ghép.")
    doc.add_heading("6. Kết quả khảo sát labevents", level=1)
    add_table(doc, ["Chỉ số", "Số dòng", "Tỷ lệ"], LAB_FUNNEL, [9.0, 4.2, 3.2], 8)
    add_bullets(doc, ["79.439.965 dòng thuộc 447.165 admission và 817 itemid đạt định nghĩa nền tảng sử dụng của Task 3.", "73.768.897 dòng thiếu hadm_id, chiếm 46,579%, không thể ghép chắc chắn với nhãn theo admission.", "Có 10 itemid mang nhiều chuỗi đơn vị; chưa chuyển đổi hoặc gộp.", "Trung vị số dòng xét nghiệm trên admission là 85 khi chưa giới hạn cửa sổ sớm.", "Tần suất itemid chỉ phản ánh độ phủ, chưa phải quyết định chọn feature."])
    doc.add_heading("7. Kết quả khảo sát diagnoses_icd", level=1)
    add_bullets(doc, ["Toàn bộ 6.364.488 dòng đủ khóa, khớp admission, bệnh nhân và d_icd_diagnoses.", "Có 545.497 admission; 2.908.741 dòng ICD-9 và 3.455.747 dòng ICD-10.", "Có 9.143 mã ICD-9, 19.440 mã ICD-10 và 28.583 cặp icd_code + icd_version.", "Trung vị 10 diagnosis/admission; tối đa 57 dòng diagnosis.", "Có 388 dòng dư trùng mã trong admission và 10 admission có nhiều hơn một seq_num = 1.", "Dữ liệu đủ nền tảng để ánh xạ PheCode, nhưng chưa phải label set cuối."])
    doc.add_heading("8. Thời gian và ba kịch bản cửa sổ", level=1)
    doc.add_paragraph("Quy ước khảo sát: 0 ≤ charttime − admittime < độ dài cửa sổ, chỉ trên admission có interval không âm.")
    add_table(doc, ["Cửa sổ", "Dòng lab", "Admission có lab", "Lab + diagnosis", "Tỷ lệ toàn admission"], WINDOWS, [2.3, 3.1, 3.4, 3.5, 4.1], 7.5)
    add_bullets(doc, ["175/546.028 admission có dischtime trước admittime.", "Trong 79.439.965 dòng ứng viên: 2.816.594 dòng trước admittime, 76.563.834 dòng trong admission, 53.761 dòng sau dischtime và 5.776 dòng thuộc interval âm hoặc thiếu.", "Thời điểm lab đầu tiên có trung vị 7,12 giờ sau admittime trong các admission có lab nằm trong admission.", "Cửa sổ dài hơn tăng độ phủ nhưng có thể tăng nguy cơ sử dụng thông tin muộn; Phase 1 chưa lựa chọn cửa sổ."])
    doc.add_heading("9. Danh mục vấn đề chất lượng", level=1)
    add_table(doc, ["Nguồn", "Vấn đề", "Mức ảnh hưởng", "Hướng xử lý/điểm cần chốt"], QUALITY, [2.7, 4.0, 3.8, 7.0], 7.2)
    doc.add_paragraph("Mọi thao tác loại hoặc biến đổi sẽ được thực hiện trên dataset xử lý mới và phải ghi lại số dòng/admission trước và sau; dữ liệu gốc tiếp tục được giữ nguyên.")
    doc.add_heading("10. Kiểm soát nguy cơ leakage", level=1)
    add_table(doc, ["Nguy cơ", "Kiểm soát bắt buộc"], LEAKAGE, [6.0, 11.5], 8)
    doc.add_heading("11. Quyết định và vấn đề còn mở", level=1)
    add_table(doc, ["Mã", "Trạng thái", "Chủ đề", "Nội dung"], DECISIONS, [1.2, 3.0, 3.5, 9.8], 7.5)
    doc.add_heading("12. Đánh giá kết thúc Phase 1", level=1)
    doc.add_paragraph("Phase 1 đã hoàn thành mục tiêu khảo sát cấu trúc, độ đầy đủ, liên kết, thời gian và rủi ro dữ liệu. Dữ liệu hiện chưa phải dataset sạch để huấn luyện vì cohort và các quy tắc biến đổi chưa được chốt. Tuy nhiên, bằng chứng hiện có đủ để bắt đầu Phase 2 ở mức đặc tả cohort và so sánh kịch bản.")
    doc.add_heading("13. Đề xuất kế hoạch Phase 2 — Cohort Definition", level=1)
    p = doc.add_paragraph()
    r = p.add_run("Trạng thái: kế hoạch đề xuất để thảo luận, chưa phải quyết định nghiên cứu đã chốt.")
    r.bold = True
    r.font.color.rgb = RGBColor(156, 87, 0)
    add_table(doc, ["Task", "Nội dung", "Công việc", "Đầu ra", "Trạng thái"], PHASE2_PLAN, [1.3, 3.0, 5.7, 4.8, 2.5], 7.0)
    doc.add_heading("14. Các quyết định cần thống nhất trước khi tạo cohort", level=1)
    add_numbered(doc, ["Mục tiêu cân bằng giữa độ phủ và tính sớm khi chọn cửa sổ 6, 12 hoặc 24 giờ.", "Cách xử lý admission có dischtime < admittime và admission kết thúc trước độ dài cửa sổ.", "Có yêu cầu tối thiểu về loại xét nghiệm hoặc số kết quả trong cửa sổ hay không.", "Cách sử dụng storetime khi kết quả được lưu sau mốc dự đoán.", "Có giới hạn quần thể bổ sung theo tuổi, loại admission hoặc đặc điểm khác hay không; chỉ bổ sung nếu phục vụ câu hỏi nghiên cứu."])
    doc.add_heading("15. Kết luận báo cáo định kỳ", level=1)
    doc.add_paragraph("Giai đoạn 1 đã tạo nền tảng có thể tái lập cho việc định nghĩa cohort: dữ liệu nguồn được giữ nguyên, các bảng và khóa chính đã được xác minh, độ phủ của các kịch bản thời gian đã được đo, và các nguy cơ leakage đã được ghi thành quy tắc kiểm soát. Công việc tiếp theo nên bắt đầu bằng việc duyệt đặc tả cohort, sau đó mới tạo tập dữ liệu xử lý mới cho Phase 2.")

    doc.add_heading("Phụ lục A. Sản phẩm đã hoàn thành", level=1)
    add_table(doc, ["Nhóm sản phẩm", "Thành phần", "Mục đích"], DELIVERABLES, [4.0, 7.0, 6.5], 7.5)
    doc.add_heading("Phụ lục B. Nội dung đề nghị trao đổi với giảng viên", level=1)
    add_bullets(doc, [
        "Xác nhận phạm vi bài toán multi-label ở cấp admission và cách diễn giải đầu ra.",
        "Thảo luận tiêu chí chọn cửa sổ 6/12/24 giờ dựa trên độ phủ và nguy cơ thông tin muộn.",
        "Thống nhất inclusion/exclusion criteria trước khi tạo cohort xử lý.",
        "Xác nhận lộ trình: Phase 2 cohort, Phase 3 lab features, Phase 4 PheCode labels, sau đó mới ghép dataset và modeling.",
        "Ghi nhận rằng Phase 1 mới khảo sát và kiểm soát rủi ro; chưa có kết quả mô hình.",
    ])

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run("Báo cáo định kỳ Phase 1 — MIMIC-IV HOSP")
    rr.font.name = "Arial"
    rr.font.size = Pt(8)
    rr.font.color.rgb = RGBColor(100, 100, 100)
    props = doc.core_properties
    props.title = "REPORT_task07 — Tổng kết Giai đoạn 1"
    props.subject = "Báo cáo định kỳ khảo sát dữ liệu MIMIC-IV HOSP"
    props.author = "Nhóm nghiên cứu"
    props.keywords = "MIMIC-IV, HOSP, laboratory, multi-label, Phase 1, cohort"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_csv(DECISION_PATH, ["decision_id", "status", "topic", "decision_or_open_question"], DECISIONS)
    write_csv(PLAN_PATH, ["task", "name", "work", "deliverable", "status"], PHASE2_PLAN)
    write_markdown()
    write_docx()
    for path in [MD_PATH, DOCX_PATH, DECISION_PATH, PLAN_PATH]:
        print(path)
