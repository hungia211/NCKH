"""Build Markdown and Word reports for a completed Task 3 run."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "AGENTS.md").is_file())


def fmt(value: int | float) -> str:
    return f"{value:,.0f}"


def pct(value: int | float, total: int | float) -> str:
    return f"{100 * value / total:.3f}%"


def markdown_table(headers: list[str], rows: list[list[object]]) -> list[str]:
    result = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    result.extend("| " + " | ".join(str(cell) for cell in row) + " |" for row in rows)
    return result


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
        for run in cell.paragraphs[0].runs: run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row): cell.text = str(value)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths): cell.width = Cm(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def build_docx(path: Path, sections: list[tuple]) -> None:
    template = ROOT / "tmp" / "phase1_reference.docx"
    doc = Document(template) if template.exists() else Document()
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"): body.remove(child)
    for section in doc.sections:
        section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0); section.left_margin = Cm(2.5); section.right_margin = Cm(2.0)
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            for paragraph in part.paragraphs: paragraph.clear()
    for name, size in (("Normal", 12), ("Title", 18), ("Heading 1", 15), ("Heading 2", 13)):
        style = doc.styles[name]; style.font.name = "Times New Roman"; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6); style.paragraph_format.line_spacing = 1.25
    doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for entry in sections:
        kind = entry[0]
        if kind == "title": doc.add_paragraph(entry[1], style="Title")
        elif kind == "subtitle": doc.add_paragraph(entry[1])
        elif kind == "h1": doc.add_paragraph(entry[1], style="Heading 1")
        elif kind == "h2": doc.add_paragraph(entry[1], style="Heading 2")
        elif kind == "p": doc.add_paragraph(entry[1])
        elif kind == "bullet":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.add_run("• " + entry[1])
        elif kind == "table": add_table(doc, entry[1], entry[2], entry[3] if len(entry) > 3 else None)
    doc.core_properties.title = "Báo cáo khảo sát labevents MIMIC IV"
    doc.core_properties.author = ""
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("run_dir", type=Path); args = parser.parse_args()
    run = args.run_dir.resolve(); data = json.loads((run / "results.json").read_text(encoding="utf-8")); c = data["counts"]; total = c["rows"]
    items = pd.read_csv(run / "lab_item_profile.csv")
    units = pd.read_csv(run / "lab_unit_profile.csv")
    top = items.head(15)
    multi_ids = set(items.loc[items["distinct_units"] > 1, "itemid"].astype(int))
    multi = units.loc[units["itemid"].isin(multi_ids)].copy()
    zero = items.loc[items["candidate_rows"] == 0].sort_values("rows", ascending=False).head(12)
    q = data["candidate_measurements_per_admission"]

    top_rows = [[int(r.itemid), r.label, fmt(r.candidate_admissions), fmt(r.candidate_rows), f"{100*r.candidate_row_rate:.2f}%"] for r in top.itertuples()]
    multi_rows = [[int(itemid), items.loc[items.itemid.eq(itemid), "label"].iloc[0], unit, fmt(rows)] for itemid, unit, rows in multi[["itemid", "valueuom", "rows"]].itertuples(index=False, name=None)]
    zero_rows = [[int(r.itemid), r.label, fmt(r.rows), fmt(r.no_result_rows)] for r in zero.itertuples()]
    funnel_rows = [
        ["Tổng số dòng", fmt(total), "100.000%"],
        ["Có hadm_id", fmt(total - c["missing_hadm_id"]), pct(total - c["missing_hadm_id"], total)],
        ["Khớp admission và subject_id", fmt(c["subject_match_after_known_admission"]), pct(c["subject_match_after_known_admission"], total)],
        ["Có kết quả số", fmt(c["numeric_result"]), pct(c["numeric_result"], total)],
        ["Chỉ có kết quả chữ", fmt(c["text_only_result"]), pct(c["text_only_result"], total)],
        ["Không có kết quả số hoặc chữ", fmt(c["no_numeric_or_text_result"]), pct(c["no_numeric_or_text_result"], total)],
        ["Đủ điều kiện nền tảng sử dụng", fmt(c["candidate_usable_row"]), pct(c["candidate_usable_row"], total)],
    ]

    md = [
        "# Báo cáo Task 3 khảo sát labevents", "",
        "## 1 Mục đích và phạm vi", "",
        "Task 3 khảo sát toàn bộ bảng `labevents` để xác định dòng nào có đủ nền tảng liên kết và có kết quả có thể phân tích. Task không sửa dữ liệu nguồn, không sử dụng bảng reference, không chọn xét nghiệm, không chuẩn hóa đơn vị và không áp dụng cửa sổ thời gian.", "",
        "Một dòng được ghi nhận là có nền tảng sử dụng khi có `subject_id`, `hadm_id`, `itemid`; `hadm_id` tồn tại trong `admissions`; `subject_id` khớp chủ sở hữu admission; `itemid` có trong `d_labitems`; `charttime` đọc được; và có `valuenum` hoặc `value` không rỗng. Đây là chỉ số khảo sát, chưa phải quy tắc loại dữ liệu cuối cùng.", "",
        "## 2 Dữ liệu và cách thực hiện", "",
        f"Script đã đọc toàn bộ {fmt(total)} dòng `labevents` theo khối 500.000 dòng. Các phép đếm được tổng hợp theo dòng, `itemid` và `hadm_id`. Không lấy mẫu. File nguồn không bị ghi đè.", "",
        "## 3 Khả năng sử dụng theo dòng", "",
        *markdown_table(["Chỉ số", "Số dòng", "Tỷ lệ trên toàn bảng"], funnel_rows), "",
        f"Có {fmt(c['missing_hadm_id'])} dòng thiếu `hadm_id`, tương đương {pct(c['missing_hadm_id'], total)}. Các dòng này chưa thể gắn trực tiếp vào đơn vị phân tích `hadm_id`. Trong các dòng có đủ khóa, {fmt(c['subject_mismatch_after_known_admission'])} dòng có `subject_id` khác với chủ sở hữu admission.", "",
        f"Tổng cộng {fmt(c['candidate_usable_row'])} dòng đáp ứng định nghĩa nền tảng sử dụng, gồm {fmt(c['candidate_numeric_row'])} dòng có `valuenum` và {fmt(c['candidate_text_only_row'])} dòng chỉ có kết quả chữ. Chúng thuộc {fmt(data['candidate_admissions'])} admission và {fmt(data['candidate_itemids'])} itemid.", "",
        "## 4 Thời gian và trường hỗ trợ", "",
        f"Không có dòng nào thiếu hoặc không đọc được `charttime`. Có {fmt(c['missing_storetime'])} dòng thiếu `storetime` ({pct(c['missing_storetime'], total)}). Task này chỉ kiểm tra sự hiện diện của thời gian; quan hệ với `admittime` và `dischtime` thuộc Task 5.", "",
        f"Có {fmt(c['missing_valueuom'])} dòng thiếu `valueuom` ({pct(c['missing_valueuom'], total)}). Có {fmt(c['has_both_reference_bounds'])} dòng có đủ cả hai cận tham chiếu. Các cột reference chưa được dùng để phân loại kết quả trong task này.", "",
        "## 5 Độ phủ theo xét nghiệm", "",
        f"Bảng có {fmt(data['unique_itemid_observed'])} itemid xuất hiện. Trong đó {fmt(data['candidate_itemids'])} itemid có ít nhất một dòng đáp ứng định nghĩa nền tảng sử dụng; {fmt(data['unique_itemid_observed'] - data['candidate_itemids'])} itemid không có dòng nào đáp ứng.", "",
        *markdown_table(["itemid", "Xét nghiệm", "Admission", "Dòng sử dụng", "Tỷ lệ dòng sử dụng"], top_rows), "",
        "Tần suất cao chỉ phản ánh độ phủ. Danh sách trên chưa phải danh sách feature được chọn.", "",
        "## 6 Xét nghiệm không có dòng sử dụng", "",
        f"Có {fmt(data['unique_itemid_observed'] - data['candidate_itemids'])} itemid không có dòng đáp ứng định nghĩa. Nhiều mục trong nhóm này có dòng sự kiện nhưng trường `value` và `valuenum` trống. Các mục có nhiều dòng nhất được liệt kê dưới đây.", "",
        *markdown_table(["itemid", "Xét nghiệm", "Tổng dòng", "Không có kết quả"], zero_rows), "",
        "## 7 Đơn vị đo", "",
        f"Có {fmt(data['items_with_multiple_observed_units'])} itemid xuất hiện với hơn một chuỗi đơn vị không rỗng. Task chỉ ghi nhận chuỗi đơn vị, chưa kết luận các đơn vị tương đương và chưa chuyển đổi.", "",
        *markdown_table(["itemid", "Xét nghiệm", "Đơn vị", "Số dòng"], multi_rows), "",
        "## 8 Số lần đo trên admission", "",
        f"Trong các admission có dòng sử dụng, trung vị là {fmt(q['median'])} dòng xét nghiệm; phân vị 25 là {fmt(q['p25'])}, phân vị 75 là {fmt(q['p75'])}, phân vị 95 là {fmt(q['p95'])}, và lớn nhất là {fmt(q['max'])}. Các con số tính trên toàn admission, chưa giới hạn giai đoạn sớm nên chưa đại diện cho số feature cuối cùng.", "",
        "## 9 Kết luận của Task 3", "",
        f"Bảng `labevents` có {fmt(data['candidate_admissions'])} admission và {fmt(data['candidate_itemids'])} itemid có nền tảng sử dụng theo định nghĩa khảo sát. Vấn đề lớn nhất ở cấp dòng là thiếu `hadm_id`; ngoài ra cần xử lý riêng kết quả chữ, dòng không có kết quả và 10 itemid có nhiều đơn vị. Chưa có dòng nào bị xóa hoặc sửa.", "",
        "Bước tiếp theo cần khảo sát thời gian lab so với admission trong Task 5 trước khi quyết định cohort hoặc cửa sổ dữ liệu.", "",
        "## 10 File đầu ra", "",
        "- `lab_item_profile.csv`: thống kê toàn bộ itemid.",
        "- `lab_unit_profile.csv`: tần suất từng chuỗi đơn vị theo itemid.",
        "- `row_usability_funnel.csv`: các điều kiện đầy đủ và liên kết theo dòng.",
        "- `measurements_per_admission.csv`: phân bố số dòng sử dụng trên admission.",
        "- `results.json`: metadata và số liệu máy đọc.",
        "- `task03_labevents_review.xlsx`: workbook dành cho người đọc.",
    ]
    (run / "REPORT_task03.md").write_text("\n".join(md), encoding="utf-8")

    sections = [("title", "BÁO CÁO KHẢO SÁT LABEVENTS MIMIC IV"), ("subtitle", "Task 3 của Giai đoạn 1 khảo sát khả năng sử dụng dữ liệu xét nghiệm")]
    sections += [("h1", "1 Mục đích và phạm vi"),
                 ("p", "Task 3 khảo sát toàn bộ bảng labevents để xác định dòng nào có đủ nền tảng liên kết và có kết quả có thể phân tích. Task không sửa dữ liệu nguồn, không sử dụng bảng reference, không chọn xét nghiệm, không chuẩn hóa đơn vị và không áp dụng cửa sổ thời gian."),
                 ("p", "Một dòng được ghi nhận là có nền tảng sử dụng khi có subject_id, hadm_id và itemid; hadm_id tồn tại trong admissions; subject_id khớp chủ sở hữu admission; itemid có trong d_labitems; charttime đọc được; và có valuenum hoặc value không rỗng. Đây là chỉ số khảo sát, chưa phải quy tắc loại dữ liệu cuối cùng."),
                 ("h1", "2 Dữ liệu và cách thực hiện"),
                 ("p", f"Script đã đọc toàn bộ {fmt(total)} dòng labevents theo khối 500.000 dòng. Các phép đếm được tổng hợp theo dòng, itemid và hadm_id. Không lấy mẫu. File nguồn không bị ghi đè."),
                 ("h1", "3 Khả năng sử dụng theo dòng"), ("table", ["Chỉ số", "Số dòng", "Tỷ lệ"], funnel_rows, [8.5, 3.5, 3.5]),
                 ("p", f"Có {fmt(c['missing_hadm_id'])} dòng thiếu hadm_id, tương đương {pct(c['missing_hadm_id'], total)}. Các dòng này chưa thể gắn trực tiếp vào đơn vị phân tích hadm_id. Trong các dòng có đủ khóa, {fmt(c['subject_mismatch_after_known_admission'])} dòng có subject_id khác với chủ sở hữu admission."),
                 ("p", f"Tổng cộng {fmt(c['candidate_usable_row'])} dòng đáp ứng định nghĩa nền tảng sử dụng, gồm {fmt(c['candidate_numeric_row'])} dòng có valuenum và {fmt(c['candidate_text_only_row'])} dòng chỉ có kết quả chữ. Chúng thuộc {fmt(data['candidate_admissions'])} admission và {fmt(data['candidate_itemids'])} itemid."),
                 ("h1", "4 Thời gian và trường hỗ trợ"),
                 ("p", f"Không có dòng nào thiếu hoặc không đọc được charttime. Có {fmt(c['missing_storetime'])} dòng thiếu storetime ({pct(c['missing_storetime'], total)}). Quan hệ thời gian với admittime và dischtime thuộc Task 5."),
                 ("p", f"Có {fmt(c['missing_valueuom'])} dòng thiếu valueuom ({pct(c['missing_valueuom'], total)}). Có {fmt(c['has_both_reference_bounds'])} dòng có đủ cả hai cận tham chiếu. Các cột reference chưa được dùng để phân loại kết quả."),
                 ("h1", "5 Độ phủ theo xét nghiệm"),
                 ("p", f"Bảng có {fmt(data['unique_itemid_observed'])} itemid xuất hiện. Trong đó {fmt(data['candidate_itemids'])} itemid có ít nhất một dòng đáp ứng định nghĩa nền tảng sử dụng; {fmt(data['unique_itemid_observed'] - data['candidate_itemids'])} itemid không có dòng nào đáp ứng."),
                 ("table", ["itemid", "Xét nghiệm", "Admission", "Dòng sử dụng", "Tỷ lệ"], top_rows, [2, 6.5, 2.5, 2.8, 2]),
                 ("p", "Tần suất cao chỉ phản ánh độ phủ. Danh sách này chưa phải danh sách feature được chọn."),
                 ("h1", "6 Xét nghiệm không có dòng sử dụng"), ("p", f"Có {fmt(data['unique_itemid_observed'] - data['candidate_itemids'])} itemid không có dòng đáp ứng định nghĩa. Nhiều mục có dòng sự kiện nhưng cả value và valuenum đều trống."),
                 ("table", ["itemid", "Xét nghiệm", "Tổng dòng", "Không có kết quả"], zero_rows, [2, 7, 3, 3]),
                 ("h1", "7 Đơn vị đo"), ("p", f"Có {fmt(data['items_with_multiple_observed_units'])} itemid có hơn một chuỗi đơn vị không rỗng. Chưa thực hiện chuyển đổi đơn vị."),
                 ("table", ["itemid", "Xét nghiệm", "Đơn vị", "Số dòng"], multi_rows, [2, 6.5, 3.5, 3]),
                 ("h1", "8 Số lần đo trên admission"), ("p", f"Trung vị là {fmt(q['median'])} dòng; phân vị 25 là {fmt(q['p25'])}, phân vị 75 là {fmt(q['p75'])}, phân vị 95 là {fmt(q['p95'])}, và lớn nhất là {fmt(q['max'])}. Số liệu chưa giới hạn giai đoạn sớm."),
                 ("h1", "9 Kết luận"),
                 ("p", f"Bảng labevents có {fmt(data['candidate_admissions'])} admission và {fmt(data['candidate_itemids'])} itemid có nền tảng sử dụng theo định nghĩa khảo sát. Vấn đề lớn nhất ở cấp dòng là thiếu hadm_id; ngoài ra cần xử lý riêng kết quả chữ, dòng không có kết quả và 10 itemid có nhiều đơn vị. Chưa có dòng nào bị xóa hoặc sửa."),
                 ("p", "Bước tiếp theo cần khảo sát thời gian lab so với admission trong Task 5 trước khi quyết định cohort hoặc cửa sổ dữ liệu."),
                 ("h1", "10 File đầu ra"),
                 ("p", "Kết quả gồm lab_item_profile.csv, lab_unit_profile.csv, row_usability_funnel.csv, measurements_per_admission.csv, results.json và task03_labevents_review.xlsx.")]
    build_docx(run / "REPORT_task03.docx", sections)
    print(run / "REPORT_task03.docx")


if __name__ == "__main__": main()
