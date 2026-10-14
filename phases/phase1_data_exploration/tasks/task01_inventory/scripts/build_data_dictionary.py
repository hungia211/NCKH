from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TABLE_DIR = ROOT / "outputs" / "tables"
REPORT_DIR = ROOT / "report"
CSV_PATH = TABLE_DIR / "data_dictionary_core_tables_task01.csv"
JSON_PATH = TABLE_DIR / "data_dictionary_core_tables_task01.json"
MD_PATH = REPORT_DIR / "DATA_DICTIONARY_task01.md"
DOCX_PATH = REPORT_DIR / "DATA_DICTIONARY_task01.docx"

SOURCES = {
    "admissions": "https://mimic.mit.edu/docs/iv/modules/hosp/admissions.html",
    "labevents": "https://mimic.mit.edu/docs/iv/modules/hosp/labevents.html",
    "d_labitems": "https://mimic.mit.edu/docs/iv/modules/hosp/d_labitems.html",
    "diagnoses_icd": "https://mimic.mit.edu/docs/iv/modules/hosp/diagnoses_icd.html",
    "d_icd_diagnoses": "https://mimic.mit.edu/docs/iv/modules/hosp/d_icd_diagnoses.html",
}

FIELDS = [
    "table", "column_position", "column_name", "postgres_type", "nullable",
    "meaning_vi", "example_or_values", "research_role", "join_rule",
    "use_in_current_pipeline", "local_observation", "official_source",
]


def row(table, pos, name, dtype, nullable, meaning, examples, role, join, use, local=""):
    return {
        "table": table,
        "column_position": pos,
        "column_name": name,
        "postgres_type": dtype,
        "nullable": nullable,
        "meaning_vi": meaning,
        "example_or_values": examples,
        "research_role": role,
        "join_rule": join,
        "use_in_current_pipeline": use,
        "local_observation": local,
        "official_source": SOURCES[table],
    }


ROWS = [
    # admissions
    row("admissions", 1, "subject_id", "INTEGER NOT NULL", "Không", "Mã định danh duy nhất của bệnh nhân.", "Số nguyên", "Khóa bệnh nhân; khóa chia tập", "Nối với subject_id ở labevents và diagnoses_icd.", "Có", "223.452 bệnh nhân trong dữ liệu khảo sát."),
    row("admissions", 2, "hadm_id", "INTEGER NOT NULL", "Không", "Mã định danh duy nhất của một lần nhập viện.", "Số nguyên", "Đơn vị phân tích chính", "Nối với hadm_id ở labevents và diagnoses_icd.", "Có", "546.028 admission; là khóa trung tâm của bài toán."),
    row("admissions", 3, "admittime", "TIMESTAMP NOT NULL", "Không", "Thời điểm bệnh nhân được ghi nhận nhập viện.", "YYYY-MM-DD hh:mm:ss", "Mốc bắt đầu admission", "So sánh với labevents.charttime để chọn xét nghiệm sau nhập viện.", "Có", "Thời gian đã dịch chuyển nhưng thứ tự tương đối trong cùng bệnh nhân vẫn dùng được."),
    row("admissions", 4, "dischtime", "TIMESTAMP", "Có thể", "Thời điểm kết thúc lần nằm viện hoặc xuất viện.", "YYYY-MM-DD hh:mm:ss", "Mốc kết thúc admission", "Dùng kiểm tra charttime có nằm trong admission.", "Có", "Task 5 ghi nhận 175 admission có dischtime < admittime; cần loại/điều tra khi tạo cohort."),
    row("admissions", 5, "deathtime", "TIMESTAMP", "Có", "Thời điểm tử vong trong bệnh viện nếu có.", "Timestamp hoặc trống", "Kết cục và kiểm tra dữ liệu", "Gắn trực tiếp với hadm_id trong cùng dòng admissions.", "Không trực tiếp", "Không phải nhãn phenotype hiện tại."),
    row("admissions", 6, "admission_type", "VARCHAR(40) NOT NULL", "Không", "Loại hoặc mức độ khẩn cấp của lần nhập viện.", "ELECTIVE, URGENT, EW EMER., OBSERVATION...", "Thông tin bối cảnh", "Không cần nối.", "Cần quyết định", "Có thể dùng mô tả cohort hoặc phân tầng sau này."),
    row("admissions", 7, "admit_provider_id", "VARCHAR(10)", "Có", "Mã ẩn danh của nhân viên y tế phụ trách tiếp nhận.", "Chuỗi mã", "Metadata", "Có thể nối bảng provider nếu cần.", "Không trực tiếp", "Mã chỉ có ý nghĩa định danh nhất quán, không mang ý nghĩa lâm sàng nội tại."),
    row("admissions", 8, "admission_location", "VARCHAR(60)", "Có", "Nơi bệnh nhân ở ngay trước khi nhập viện.", "EMERGENCY ROOM, TRANSFER FROM HOSPITAL...", "Thông tin bối cảnh", "Không cần nối.", "Cần quyết định", "Dùng mô tả nguồn nhập viện nếu cần."),
    row("admissions", 9, "discharge_location", "VARCHAR(60)", "Có", "Nơi hoặc hình thức chuyển đến sau khi xuất viện.", "HOME, SKILLED NURSING FACILITY...", "Thông tin kết cục", "Không cần nối.", "Không trực tiếp", "Thông tin xảy ra sau admission, tránh dùng làm feature đầu vào sớm."),
    row("admissions", 10, "insurance", "VARCHAR(255)", "Có", "Loại bảo hiểm được ghi nhận cho admission.", "Medicare, Medicaid, Other...", "Nhân khẩu học/bối cảnh", "Không cần nối.", "Cần quyết định", "Có thể thay đổi giữa các admission của cùng bệnh nhân."),
    row("admissions", 11, "language", "VARCHAR(10)", "Có", "Ngôn ngữ được ghi nhận của bệnh nhân.", "ENGLISH hoặc mã khác", "Nhân khẩu học/bối cảnh", "Không cần nối.", "Cần quyết định", "Không phải feature xét nghiệm."),
    row("admissions", 12, "marital_status", "VARCHAR(30)", "Có", "Tình trạng hôn nhân được ghi nhận tại admission.", "MARRIED, SINGLE...", "Nhân khẩu học/bối cảnh", "Không cần nối.", "Cần quyết định", "Có thể thay đổi giữa các admission."),
    row("admissions", 13, "race", "VARCHAR(80)", "Có", "Nhóm chủng tộc/sắc tộc được ghi nhận.", "Chuỗi phân loại", "Nhân khẩu học/bối cảnh", "Không cần nối.", "Cần quyết định", "Cần phân tích công bằng nếu đưa vào mô hình."),
    row("admissions", 14, "edregtime", "TIMESTAMP", "Có", "Thời điểm đăng ký tại khoa cấp cứu.", "Timestamp hoặc trống", "Mốc thời gian ED", "So với admittime khi admission đi qua ED.", "Không trực tiếp", "Có thể sớm hơn admittime; không phải mọi admission đều qua ED."),
    row("admissions", 15, "edouttime", "TIMESTAMP", "Có", "Thời điểm rời khoa cấp cứu.", "Timestamp hoặc trống", "Mốc thời gian ED", "So với admittime và charttime khi cần.", "Không trực tiếp", "Không phải mọi admission đều qua ED."),
    row("admissions", 16, "hospital_expire_flag", "SMALLINT", "Có thể", "Cờ tử vong trong lần nằm viện: 1 là tử vong, 0 là sống khi xuất viện.", "0 hoặc 1", "Kết cục và kiểm tra dữ liệu", "Gắn trực tiếp với hadm_id.", "Không trực tiếp", "Không phải nhãn phenotype theo đề cương hiện tại."),

    # labevents
    row("labevents", 1, "labevent_id", "INTEGER NOT NULL", "Không", "Mã duy nhất của một dòng kết quả xét nghiệm.", "Số nguyên", "Khóa dòng", "Không dùng để nối cohort.", "Không trực tiếp", "Dùng truy vết và loại trùng nếu cần."),
    row("labevents", 2, "subject_id", "INTEGER NOT NULL", "Không", "Mã bệnh nhân sở hữu kết quả xét nghiệm.", "Số nguyên", "Khóa bệnh nhân; kiểm tra nối", "Phải khớp subject_id của hadm_id trong admissions.", "Có", "Dùng kiểm tra toàn vẹn và chia tập theo bệnh nhân."),
    row("labevents", 3, "hadm_id", "INTEGER", "Có", "Mã admission mà kết quả xét nghiệm được gán vào.", "Số nguyên hoặc trống", "Khóa admission", "Nối admissions.hadm_id và diagnoses_icd.hadm_id.", "Có; bắt buộc", "73.768.897/158.374.764 dòng (46,579%) thiếu hadm_id và không thuộc tập ứng viên train theo admission."),
    row("labevents", 4, "specimen_id", "INTEGER NOT NULL", "Không", "Mã mẫu bệnh phẩm; nhiều xét nghiệm có thể dùng chung một mẫu.", "Số nguyên", "Nhóm mẫu bệnh phẩm", "Nhóm các dòng cùng mẫu khi phân tích kỹ thuật.", "Cần quyết định", "Không thay thế hadm_id."),
    row("labevents", 5, "itemid", "INTEGER NOT NULL", "Không", "Mã khái niệm xét nghiệm.", "Số nguyên", "Khóa feature xét nghiệm", "Nối d_labitems.itemid để lấy tên, loại dịch và nhóm xét nghiệm.", "Có", "976 itemid thực sự xuất hiện trong labevents."),
    row("labevents", 6, "order_provider_id", "VARCHAR(10)", "Có", "Mã ẩn danh của nhân viên y tế đặt xét nghiệm.", "Chuỗi mã", "Metadata", "Có thể nối provider nếu cần.", "Không trực tiếp", "Không phải feature xét nghiệm chính."),
    row("labevents", 7, "charttime", "TIMESTAMP(0)", "Có thể", "Thời điểm kết quả được ghi nhận; thường gần thời điểm lấy mẫu.", "YYYY-MM-DD hh:mm:ss", "Mốc thời gian xét nghiệm", "So với admissions.admittime và dischtime để chọn đúng khoảng thời gian.", "Có; dùng để lọc", "Không nhất thiết là lúc bác sĩ nhìn thấy kết quả."),
    row("labevents", 8, "storetime", "TIMESTAMP(0)", "Có thể", "Thời điểm kết quả được lưu và có thể truy cập trong hệ thống xét nghiệm.", "YYYY-MM-DD hh:mm:ss", "Mốc sẵn có của kết quả", "So với charttime và mốc dự đoán để kiểm soát rò rỉ tương lai.", "Có; dùng kiểm tra", "Task 6 dùng trường này để đánh giá nguy cơ kết quả có sau mốc dự đoán."),
    row("labevents", 9, "value", "VARCHAR(200)", "Có", "Giá trị kết quả ở dạng gốc, có thể là số hoặc văn bản.", "7.4, NEGATIVE, ___...", "Giá trị thô", "Kết hợp itemid và valueuom để diễn giải.", "Cần quyết định", "Không ép toàn bộ thành số; một số xét nghiệm mang giá trị định tính."),
    row("labevents", 10, "valuenum", "DOUBLE PRECISION", "Có", "Phiên bản số của value khi kết quả có thể chuyển thành số.", "Số thực hoặc trống", "Feature số ứng viên", "Kết hợp itemid, valueuom và mốc thời gian.", "Có với xét nghiệm số", "Dòng định tính có thể không có valuenum nhưng vẫn có value."),
    row("labevents", 11, "valueuom", "VARCHAR(20)", "Có", "Đơn vị đo của kết quả xét nghiệm.", "mg/dL, mmol/L...", "Chuẩn hóa feature", "Phân tích theo cặp itemid + valueuom.", "Có", "Có 10 itemid xuất hiện với nhiều đơn vị; phải chuẩn hóa trước khi gộp giá trị."),
    row("labevents", 12, "ref_range_lower", "DOUBLE PRECISION", "Có", "Giới hạn dưới của khoảng tham chiếu.", "Số thực hoặc trống", "Thông tin tham chiếu", "Diễn giải cùng itemid và valueuom.", "Cần quyết định", "Không mặc định dùng làm feature."),
    row("labevents", 13, "ref_range_upper", "DOUBLE PRECISION", "Có", "Giới hạn trên của khoảng tham chiếu.", "Số thực hoặc trống", "Thông tin tham chiếu", "Diễn giải cùng itemid và valueuom.", "Cần quyết định", "Không mặc định dùng làm feature."),
    row("labevents", 14, "flag", "VARCHAR(10)", "Có", "Cờ cho biết kết quả bất thường theo hệ thống xét nghiệm.", "abnormal hoặc trống", "Feature phụ ứng viên", "Diễn giải cùng khoảng tham chiếu.", "Cần quyết định", "Cần kiểm tra tính nhất quán trước khi dùng."),
    row("labevents", 15, "priority", "VARCHAR(7)", "Có", "Mức ưu tiên xử lý xét nghiệm.", "ROUTINE, STAT", "Metadata quy trình", "Không cần nối.", "Cần quyết định", "Có thể phản ánh quy trình chăm sóc hơn là sinh lý."),
    row("labevents", 16, "comments", "TEXT", "Có", "Ghi chú đã khử định danh; đôi khi chứa kết quả hoặc diễn giải.", "Văn bản, ___ hoặc trống", "Văn bản phụ", "Không cần nối.", "Không trực tiếp", "Không dùng trong pipeline bảng số ban đầu."),

    # d_labitems
    row("d_labitems", 1, "itemid", "INTEGER", "Có thể theo schema", "Mã duy nhất của khái niệm xét nghiệm.", "Số nguyên", "Khóa từ điển xét nghiệm", "Nối labevents.itemid.", "Có", "Task 2 xác nhận 1.650 itemid, không thiếu và không trùng."),
    row("d_labitems", 2, "label", "VARCHAR(50)", "Có thể theo schema", "Tên dễ đọc của xét nghiệm.", "Glucose, Hemoglobin...", "Tên feature", "Tra cứu bằng itemid.", "Có", "Dùng đặt tên feature và báo cáo."),
    row("d_labitems", 3, "fluid", "VARCHAR(50)", "Có thể theo schema", "Loại bệnh phẩm được xét nghiệm.", "Blood, Urine...", "Phân biệt ngữ cảnh xét nghiệm", "Tra cứu bằng itemid.", "Có", "Giúp tránh gộp các phép đo cùng tên nhưng khác bệnh phẩm."),
    row("d_labitems", 4, "category", "VARCHAR(50)", "Có thể theo schema", "Nhóm xét nghiệm cấp cao.", "Chemistry, Hematology, Blood Gas...", "Nhóm feature", "Tra cứu bằng itemid.", "Có", "Hỗ trợ thống kê và chọn nhóm xét nghiệm."),

    # diagnoses_icd
    row("diagnoses_icd", 1, "subject_id", "INTEGER NOT NULL", "Không", "Mã bệnh nhân được gán chẩn đoán.", "Số nguyên", "Khóa bệnh nhân; kiểm tra nối", "Nối admissions.subject_id và kiểm tra theo hadm_id.", "Có", "Task 4 ghi nhận không thiếu ở các cột khóa."),
    row("diagnoses_icd", 2, "hadm_id", "INTEGER NOT NULL", "Không", "Mã admission được gán chẩn đoán ICD.", "Số nguyên", "Khóa admission của nhãn", "Nối admissions.hadm_id và labevents.hadm_id.", "Có; bắt buộc", "Là cầu nối giữa feature xét nghiệm và nhãn bệnh."),
    row("diagnoses_icd", 3, "seq_num", "INTEGER NOT NULL", "Không", "Thứ tự ưu tiên của chẩn đoán trong admission.", "1, 2, 3...", "Thứ hạng mã chẩn đoán", "Thuộc từng hadm_id.", "Cần quyết định", "Không phải thời điểm chẩn đoán; không dùng để chứng minh bệnh xuất hiện trước hay sau xét nghiệm."),
    row("diagnoses_icd", 4, "icd_code", "VARCHAR(7)", "Có thể theo schema", "Mã chẩn đoán ICD; dấu thập phân bị lược bỏ trong dữ liệu.", "4019, I10...", "Nguồn nhãn bệnh", "Nối d_icd_diagnoses bằng cả icd_code và icd_version; sau đó ánh xạ PheCode.", "Có", "Giữ kiểu chuỗi để không mất số 0 đầu mã."),
    row("diagnoses_icd", 5, "icd_version", "INTEGER", "Có thể theo schema", "Phiên bản hệ mã ICD.", "9 hoặc 10", "Phân biệt hệ mã", "Nối d_icd_diagnoses cùng icd_code; chọn đúng bảng ánh xạ PheCode.", "Có", "Task 4: 9.143 mã ICD-9 và 19.440 mã ICD-10 phân biệt theo cặp khóa."),

    # d_icd_diagnoses
    row("d_icd_diagnoses", 1, "icd_code", "CHAR(7) NOT NULL", "Không", "Mã ICD trong từ điển chẩn đoán.", "4019, I10...", "Khóa từ điển ICD", "Nối diagnoses_icd.icd_code cùng icd_version.", "Có", "Không nối chỉ bằng icd_code."),
    row("d_icd_diagnoses", 2, "icd_version", "INTEGER NOT NULL", "Không", "Phiên bản ICD tương ứng với mã.", "9 hoặc 10", "Khóa từ điển ICD", "Nối diagnoses_icd.icd_version cùng icd_code.", "Có", "Hai cột tạo khóa ghép của từ điển."),
    row("d_icd_diagnoses", 3, "long_title", "VARCHAR(255)", "Có thể", "Tên mô tả đầy đủ của chẩn đoán ICD.", "Chuỗi mô tả bệnh", "Tên nhãn để đọc và báo cáo", "Tra cứu theo icd_code + icd_version.", "Có", "Chỉ là từ điển; không chứng minh bệnh nhân có chẩn đoán nếu không có dòng trong diagnoses_icd."),
]


TABLE_SUMMARY = [
    ("admissions", "Một dòng cho một lần nhập viện", "hadm_id", "subject_id", "Xác định admission và các mốc thời gian"),
    ("labevents", "Một dòng cho một kết quả xét nghiệm", "labevent_id", "hadm_id, subject_id, itemid", "Nguồn feature xét nghiệm"),
    ("d_labitems", "Một dòng cho một khái niệm xét nghiệm", "itemid", "itemid", "Từ điển tên, bệnh phẩm và nhóm xét nghiệm"),
    ("diagnoses_icd", "Một dòng cho một mã ICD của admission", "hadm_id + seq_num", "hadm_id, subject_id, icd_code + icd_version", "Nguồn nhãn bệnh trước khi ánh xạ PheCode"),
    ("d_icd_diagnoses", "Một dòng cho một mã trong từ điển ICD", "icd_code + icd_version", "icd_code + icd_version", "Mô tả tên chẩn đoán"),
]


TIME_ROWS = [
    ("admissions.admittime", "Bắt đầu admission", "Mốc 0 để tính thời gian tương đối của xét nghiệm."),
    ("labevents.charttime", "Ghi nhận/lấy mẫu xét nghiệm", "Dùng xác định xét nghiệm xảy ra trước hay sau admittime và có nằm trong admission hay không."),
    ("labevents.storetime", "Kết quả có trong hệ thống", "Dùng kiểm tra kết quả đã sẵn có tại mốc dự đoán, hạn chế rò rỉ thông tin tương lai."),
    ("admissions.dischtime", "Kết thúc admission", "Giới hạn trên để kiểm tra sự kiện xét nghiệm thuộc đúng lần nhập viện."),
    ("admissions.edregtime", "Đăng ký khoa cấp cứu", "Mốc phụ cho admission đi qua ED; có thể trước admittime."),
    ("admissions.edouttime", "Rời khoa cấp cứu", "Mốc phụ cho admission đi qua ED; không có ở mọi admission."),
    ("admissions.deathtime", "Tử vong trong bệnh viện", "Mốc kết cục, không phải feature đầu vào sớm của bài toán hiện tại."),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=None, size=8):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_doc(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(3)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)
    for name, size in [("Heading 1", 16), ("Heading 2", 12)]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(31, 78, 121)


def add_table(doc, headers, values, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(hdr.cells[i], "1F4E79")
        if widths:
            hdr.cells[i].width = Cm(widths[i])
    for ridx, vals in enumerate(values):
        cells = table.add_row().cells
        for i, value in enumerate(vals):
            set_cell_text(cells[i], value, size=font_size)
            if ridx % 2:
                set_cell_shading(cells[i], "EAF2F8")
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()
    return table


def write_csv():
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    with CSV_PATH.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(ROWS)
    JSON_PATH.write_text(json.dumps({"columns": FIELDS, "rows": ROWS}, ensure_ascii=False, indent=2), encoding="utf-8")


def escape_md(value):
    return str(value).replace("|", "\\|").replace("\n", " ")


def write_markdown():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    out = [
        "# Từ điển dữ liệu các bảng chính MIMIC-IV HOSP — Task 1",
        "",
        "Tài liệu giải thích 44 cột thuộc năm bảng đang dùng trực tiếp để xây dựng bộ dữ liệu dự đoán multi-label disease phenotypes.",
        "",
        "## Cách nối dữ liệu",
        "",
        "- Đơn vị phân tích: `hadm_id`.",
        "- Chia tập: `subject_id`, để các admission của cùng bệnh nhân không nằm ở nhiều tập.",
        "- Feature xét nghiệm: `labevents`, nối `admissions` bằng `hadm_id` và nối `d_labitems` bằng `itemid`.",
        "- Nhãn bệnh: `diagnoses_icd`, nối `admissions` bằng `hadm_id`; tên ICD lấy từ `d_icd_diagnoses` bằng khóa ghép `icd_code + icd_version`.",
        "- Các dòng `labevents` thiếu `hadm_id` không thuộc tập ứng viên huấn luyện theo admission.",
        "",
        "## Cách đọc các mốc thời gian",
        "",
        "| Trường | Ý nghĩa | Cách dùng trong nghiên cứu |",
        "|---|---|---|",
    ]
    out.extend(f"| {escape_md(a)} | {escape_md(b)} | {escape_md(c)} |" for a, b, c in TIME_ROWS)
    out += [
        "",
        "Thời gian trong MIMIC đã được dịch chuyển để khử định danh, nhưng thứ tự và khoảng cách tương đối của các sự kiện trong cùng bệnh nhân vẫn dùng được. Với bài toán này, thời gian chủ yếu dùng để lọc đúng quan hệ nhập viện → xét nghiệm → kết quả sẵn có; chưa mặc định đưa timestamp thô vào vector feature.",
        "",
    ]
    for table_name, *_ in TABLE_SUMMARY:
        out += [
            f"## Bảng `{table_name}`",
            "",
            "| Vị trí | Cột | Kiểu | Có thể thiếu | Ý nghĩa tiếng Việt | Vai trò | Cách nối | Dùng hiện tại | Ghi chú dữ liệu cục bộ |",
            "|---:|---|---|---|---|---|---|---|---|",
        ]
        for r in ROWS:
            if r["table"] != table_name:
                continue
            vals = [r[k] for k in ["column_position", "column_name", "postgres_type", "nullable", "meaning_vi", "research_role", "join_rule", "use_in_current_pipeline", "local_observation"]]
            out.append("| " + " | ".join(escape_md(v) for v in vals) + " |")
        out += ["", f"Nguồn chính thức: {SOURCES[table_name]}", ""]
    MD_PATH.write_text("\n".join(out), encoding="utf-8")


def write_docx():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    style_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("TỪ ĐIỂN DỮ LIỆU CÁC BẢNG CHÍNH\nMIMIC-IV HOSP")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Task 1 — Inventory dữ liệu | Giai đoạn 1")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(89, 89, 89)

    p = doc.add_paragraph()
    p.add_run("Mục đích. ").bold = True
    p.add_run("Giải thích 44 cột thuộc năm bảng chính đang dùng để xây dựng bộ dữ liệu dự đoán multi-label disease phenotypes từ xét nghiệm giai đoạn sớm của một hospital admission.")

    doc.add_heading("1. Tổng quan năm bảng", level=1)
    add_table(doc, ["Bảng", "Đơn vị mỗi dòng", "Khóa chính", "Khóa nối", "Vai trò trong đề tài"], TABLE_SUMMARY, [3.0, 5.4, 3.2, 5.0, 7.3], 8)

    doc.add_heading("2. Cách nối dữ liệu trong bài toán", level=1)
    for text in [
        "Đơn vị phân tích là hadm_id: mỗi mẫu học đại diện cho một lần nhập viện.",
        "Chia tập theo subject_id để mọi admission của cùng bệnh nhân chỉ xuất hiện trong một tập.",
        "Feature xét nghiệm lấy từ labevents, nối admissions bằng hadm_id và nối d_labitems bằng itemid.",
        "Nhãn bệnh lấy từ diagnoses_icd; tên ICD tra trong d_icd_diagnoses bằng khóa ghép icd_code + icd_version; bước sau mới ánh xạ PheCode.",
        "Các dòng labevents thiếu hadm_id không thuộc tập ứng viên huấn luyện theo admission vì không thể ghép chắc chắn với nhãn bệnh.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3. Cách đọc các mốc thời gian", level=1)
    add_table(doc, ["Trường", "Ý nghĩa", "Cách dùng trong nghiên cứu"], TIME_ROWS, [5.3, 6.0, 13.0], 8)
    p = doc.add_paragraph()
    p.add_run("Lưu ý về khử định danh: ").bold = True
    p.add_run("MIMIC dịch chuyển ngày giờ, nhưng vẫn bảo toàn thứ tự và khoảng cách tương đối của các sự kiện trong cùng bệnh nhân. Vì vậy có thể xác định xét nghiệm trước hay sau nhập viện bằng phép so sánh thời gian tương đối. Timestamp thô hiện dùng để lọc đúng sự kiện và kiểm soát rò rỉ, chưa mặc định là một feature của mô hình.")

    doc.add_heading("4. Từ điển cột chi tiết", level=1)
    purpose = {x[0]: x[4] for x in TABLE_SUMMARY}
    for table_name, *_ in TABLE_SUMMARY:
        doc.add_heading(f"4.{list(purpose).index(table_name) + 1}. {table_name}", level=2)
        p = doc.add_paragraph()
        p.add_run("Vai trò: ").bold = True
        p.add_run(purpose[table_name])
        vals = []
        for r in ROWS:
            if r["table"] == table_name:
                note = r["join_rule"]
                if r["local_observation"]:
                    note += " " + r["local_observation"]
                vals.append((r["column_name"], r["postgres_type"], r["nullable"], r["meaning_vi"], r["research_role"], note))
        add_table(doc, ["Cột", "Kiểu dữ liệu", "Thiếu?", "Ý nghĩa", "Vai trò nghiên cứu", "Liên kết và ghi chú"], vals, [3.2, 3.2, 2.0, 6.1, 4.3, 7.1], 7.5)

    doc.add_heading("5. Quy tắc sử dụng hiện tại", level=1)
    rules = [
        "Giữ subject_id và hadm_id để kiểm tra toàn vẹn; không dùng các mã định danh này như feature số.",
        "Chỉ coi dòng labevents là ứng viên huấn luyện khi có hadm_id và có thể nối hợp lệ với admissions.",
        "Dùng charttime để xác định thời điểm xét nghiệm; dùng storetime để kiểm tra thời điểm kết quả thực sự có trong hệ thống.",
        "Mọi giá trị số phải được diễn giải theo itemid và valueuom; không gộp đơn vị khác nhau khi chưa chuẩn hóa.",
        "diagnoses_icd cung cấp nhãn ở cấp admission; seq_num chỉ là thứ hạng mã, không phải timestamp chẩn đoán.",
        "d_labitems và d_icd_diagnoses là bảng từ điển; bản thân một dòng từ điển không phải một quan sát của bệnh nhân.",
    ]
    for text in rules:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("6. Nguồn định nghĩa", level=1)
    for table_name, url in SOURCES.items():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{table_name}: {url}")
        r.font.size = Pt(8)

    props = doc.core_properties
    props.title = "DATA_DICTIONARY_task01"
    props.subject = "Từ điển dữ liệu các bảng chính MIMIC-IV HOSP"
    props.author = "Nhóm nghiên cứu"
    props.keywords = "MIMIC-IV, HOSP, data dictionary, admissions, labevents, ICD"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    if len(ROWS) != 44:
        raise RuntimeError(f"Expected 44 rows, got {len(ROWS)}")
    write_csv()
    write_markdown()
    write_docx()
    print(CSV_PATH)
    print(JSON_PATH)
    print(MD_PATH)
    print(DOCX_PATH)
