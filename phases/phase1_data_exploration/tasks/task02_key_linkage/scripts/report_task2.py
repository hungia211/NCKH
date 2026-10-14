"""Produce Markdown and Word reports from a completed Task 2 run."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('run_dir',type=Path); args=ap.parse_args()
    out=args.run_dir; d=json.loads((out/'results.json').read_text(encoding='utf-8'))
    c={(x['table'],x['metric']):x['count'] for x in d['checks']}
    n=lambda t,m: c[t,m]
    fmt=lambda v:f'{v:,}'
    sections=[]
    def p(s): sections.append(('p',s))
    def h(s): sections.append(('h',s))
    def table(headers,rows): sections.append(('table',[headers]+rows))
    h('1 Mục đích và phạm vi')
    p('Task 2 kiểm tra dữ liệu xét nghiệm và chẩn đoán có thể liên kết đúng với bệnh nhân và lần nhập viện hay không. Kết quả là số liệu kiểm tra khóa và độ phủ admission, phục vụ khảo sát tiếp theo. Đây chưa phải tập dữ liệu huấn luyện hoặc cohort cuối.')
    p('Input gồm sáu bảng: patients, admissions, labevents, d_labitems, diagnoses_icd và d_icd_diagnoses tại '+d['inputs'][0]['path'].rsplit('\\',1)[0]+'. Đã quét toàn bộ các dòng, đọc theo khối 500.000 dòng đối với hai bảng sự kiện. Không lấy mẫu, không chọn cửa sổ thời gian, không chọn lab hoặc ICD để loại.')
    p('Thời gian chạy UTC: '+d['started_utc']+' đến '+d['finished_utc']+'. Mỗi lần chạy lưu trong thư mục riêng; manifest chứa kích thước và thời gian sửa file nguồn, phiên bản Python/Pandas và SHA-256 script. Kiểm tra kích thước và thời gian sửa không thay thế checksum toàn bộ dữ liệu.')
    h('2 Cách kiểm tra và định nghĩa chỉ số')
    p('subject_id nhận diện bệnh nhân; hadm_id nhận diện một lần nhập viện. Một bệnh nhân có thể có nhiều admission. itemid liên kết sự kiện lab với d_labitems. Chẩn đoán được nối với d_icd_diagnoses bằng đồng thời icd_code và icd_version. ICD được đọc dạng chuỗi, giữ nguyên số 0 đầu; không tự sửa hoặc chuẩn hóa mã trong task này.')
    p('Thiếu khóa là ô rỗng trong CSV. Orphan là dòng có khóa nhưng không tìm thấy khóa ở bảng đích; không tính các dòng thiếu khóa vào orphan. Sai bệnh nhân là hadm_id tồn tại nhưng subject_id của sự kiện khác subject_id trong admissions. Trường hợp có hadm_id hợp lệ nhưng thiếu subject_id được thống kê riêng.')
    p('Trước khi nối, kiểm tra khóa của bốn bảng tham chiếu có đầy đủ và duy nhất. Duplicate key excess rows là số dòng dư sau lần xuất hiện đầu tiên của một khóa. Đây không phải kiểm tra trùng toàn bộ lab events. Nếu khóa tham chiếu thiếu hoặc trùng, script dừng trước khi nối để tránh nhân dòng không kiểm soát.')
    h('3 Tính đầy đủ và duy nhất của khóa tham chiếu')
    table(['Bảng','Khóa','Số khóa đầy đủ','Thiếu khóa','Dòng dư trùng khóa'],[
        [t,k,fmt(n(t,'distinct_complete_keys')),fmt(n(t,'missing_key')),fmt(n(t,'duplicate_key_excess_rows'))]
        for t,k in [('patients','subject_id'),('admissions','hadm_id'),('d_labitems','itemid'),('d_icd_diagnoses','icd_code + icd_version')]])
    p('admissions có '+fmt(n('admissions','missing_subject_id'))+' dòng thiếu subject_id và '+fmt(n('admissions','subject_orphan'))+' dòng subject_id không tồn tại trong patients. Khóa hadm_id duy nhất và subject_id đầy đủ cho phép dùng admissions làm bảng đối chiếu chủ sở hữu của mỗi lần nhập viện.')
    p('Có '+fmt(n('admissions','patients_with_multiple_admissions'))+' bệnh nhân có nhiều hơn một admission; nhiều nhất '+fmt(n('admissions','max_admissions_per_patient'))+' admission trên một bệnh nhân. Điều này cho thấy ghép theo subject_id đơn thuần không xác định được admission nào và cần chia train/validation/test theo bệnh nhân về sau.')
    h('4 Kết quả liên kết của hai bảng sự kiện')
    metrics=[('rows','Tổng số dòng'),('missing_subject_id','Thiếu subject_id'),('subject_orphan','subject_id không có trong patients'),('missing_hadm_id','Thiếu hadm_id'),('hadm_orphan','hadm_id không có trong admissions'),('subject_mismatch_to_admission','Sai bệnh nhân của admission'),('known_hadm_missing_subject','Có admission nhưng thiếu subject_id'),('matched_admission_and_subject','Khớp admission và bệnh nhân'),('dictionary_key_missing','Thiếu khóa dictionary'),('dictionary_orphan','Khóa không có trong dictionary')]
    table(['Phép kiểm tra','labevents','diagnoses_icd'],[[label,fmt(n('labevents',m)),fmt(n('diagnoses_icd',m))] for m,label in metrics])
    p('Mẫu số của mỗi phép kiểm tra sự kiện là tổng số dòng của chính bảng đó, được ghi trong key_checks.csv. Các phép kiểm tra khác nhau có thể cùng đánh dấu một dòng, vì vậy không cộng tất cả lỗi để suy ra số dòng bị ảnh hưởng. Riêng năm nhóm thiếu hadm_id, hadm_id orphan, sai bệnh nhân, thiếu subject_id khi có admission và khớp hoàn toàn đã được đối soát bằng tổng số dòng input.')
    mismatch_file=out/'mismatch_summary.json'
    if mismatch_file.exists():
        bad=json.loads(mismatch_file.read_text(encoding='utf-8'))
        p(f'Phát hiện quan trọng: {fmt(bad["rows"])} dòng lab bất nhất subject_id, liên quan {fmt(bad["affected_hadm_id"])} hadm_id, {fmt(bad["lab_subjects"])} subject_id ghi trong lab và {fmt(bad["admission_subjects"])} subject_id trong admission. Các bản ghi này được liệt kê trong lab_subject_mismatches.csv để truy nguyên; chưa biết nguyên nhân và chưa tự chọn bảng nào là đúng để sửa.')
    p('Số hadm_id không rỗng trong lab là '+fmt(n('labevents','unique_nonmissing_hadm_id'))+'; số có ít nhất một lab khớp cả subject_id là '+fmt(n('labevents','unique_matched_hadm_id'))+'. Chênh lệch phản ánh điều kiện kiểm tra liên kết, không phải xóa admission khỏi nguồn. Với các admission còn lab khác khớp, admission vẫn được tính là có lab, dù một số dòng của nó bất nhất.')
    h('5 Độ phủ lab và diagnosis trên admission')
    p('Mỗi admission được gán hai cờ: có ít nhất một lab khớp hadm_id và subject_id; có ít nhất một diagnosis khớp tương tự. Tất cả admission trong admissions được giữ làm mẫu số. Không yêu cầu lab có giá trị số hoặc thuộc giai đoạn sớm ở phép đo này.')
    labels={'both':'Có cả lab và diagnosis','lab_only':'Chỉ có lab','diagnosis_only':'Chỉ có diagnosis','neither':'Không có cả hai'}
    table(['Nhóm','Admission','Tỷ lệ admission','Bệnh nhân'],[[labels[g['group']],fmt(g['admissions']),f"{g['share']:.3%}",fmt(g['patients'])] for g in d['overlap']])
    p('Bốn nhóm admission không giao nhau và tổng bằng toàn bộ admissions. Số bệnh nhân giữa các nhóm không cộng được vì một bệnh nhân có thể có các admission nằm ở nhiều nhóm. Nhóm có cả hai là độ phủ liên kết ban đầu, chưa bảo đảm đủ điều kiện cohort, đủ lab trong cửa sổ sớm hoặc mapping Phecode thành công.')
    h('6 Những dòng lab thiếu hadm_id')
    missing=n('labevents','missing_hadm_id'); total=n('labevents','rows')
    p(f'Có {fmt(missing)} dòng lab thiếu hadm_id, chiếm {missing/total:.3%} tổng lab events. Có {fmt(n("labevents","patients_in_missing_hadm_rows"))} bệnh nhân được nhận diện trong nhóm này; {fmt(n("labevents","patients_in_missing_hadm_rows_with_any_admission"))} người có ít nhất một admission trong bảng admissions. Việc bệnh nhân có admission không chứng minh lab thiếu khóa thuộc admission đó.')
    p('Task 2 giữ nguyên toàn bộ dữ liệu gốc. Những dòng thiếu hadm_id được đếm riêng và không góp vào số lab liên kết trực tiếp với admission. Chưa loại vĩnh viễn, chưa tự điền hadm_id và chưa ghép theo thời gian. Nếu muốn khôi phục liên kết bằng subject_id và thời gian, cần một khảo sát riêng để đo số trường hợp không khớp, khớp duy nhất và khớp nhiều admission, rồi thống nhất quy tắc trước khi áp dụng.')
    h('7 Tránh nhân dòng khi ghép lab và diagnosis')
    p('Một admission có nhiều lab và nhiều diagnosis. Nếu nối trực tiếp hai bảng sự kiện bằng hadm_id, một admission có L lab và D diagnosis sẽ tạo L × D dòng. Với các liên kết hợp lệ đã quan sát, tổng số dòng của phép nối này dự kiến là '+fmt(d['projected_raw_lab_diagnosis_join_rows'])+'. Con số được tính từ các số đếm, không tạo bảng ghép khổng lồ trên đĩa.')
    p('Pipeline tiếp theo cần xây dựng X từ lab và Y từ diagnosis riêng, mỗi nhánh về một hàng trên hadm_id, rồi mới ghép hai nhánh. Khi bổ sung dictionary, cần bảo đảm khóa dictionary duy nhất như các kiểm tra trong Task 2.')
    h('8 Đầu ra và cách tái lập')
    p('Các file cùng thư mục báo cáo: key_checks.csv chứa từng phép kiểm tra và mẫu số; admission_overlap.csv chứa bốn nhóm độ phủ; admissions_per_patient.csv chứa phân bố số admission trên bệnh nhân; results.json chứa kết quả đầy đủ và manifest. Các CSV dùng UTF-8 BOM để mở bằng Excel. lab_subject_mismatches.csv là chi tiết khóa của các dòng bất nhất, phục vụ kiểm tra nội bộ; báo cáo Word chỉ trình bày số tổng hợp.')
    p('Chạy từ gốc repo: python phases/phase1_data_exploration/tasks/task02_key_linkage/scripts/task2_links.py. Sau khi có thư mục kết quả, chạy report_task2.py với đường dẫn thư mục đó để dựng lại Markdown và Word. Đọc dữ liệu gốc; số dòng bị xóa hoặc sửa trong nguồn là 0.')
    p('Đã đối soát tổng bốn nhóm admission, tổng các trạng thái liên kết của mỗi bảng sự kiện và kiểm tra file nguồn không thay đổi kích thước/thời gian sửa trong lúc quét. Kết quả chỉ xác nhận các quan hệ khóa được kiểm tra, không xác nhận toàn bộ dữ liệu lâm sàng là đúng.')
    h('9 Ý nghĩa đối với bước tiếp theo')
    p('Task 2 đã hoàn tất các phép kiểm tra khóa và độ phủ. Vấn đề bất nhất bệnh nhân ở một số dòng lab cần được truy nguyên trước khi chốt cohort. Những dòng này chưa được sửa hoặc loại khỏi nguồn. Task 3 tiếp theo sẽ đo tần suất theo admission, giá trị số, đơn vị và độ đầy đủ của từng lab; Task 5 sẽ đánh giá thời gian. Việc lựa chọn cohort, xử lý lab thiếu hadm_id và ngưỡng nhãn vẫn cần quyết định dựa trên các khảo sát này.')
    title='Báo cáo Task 2 kiểm tra khóa và liên kết HOSP'
    md=['# '+title,'']
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2)
    for name,size in [('Normal',12),('Title',17),('Heading 1',14)]:
        style=doc.styles[name]; style.font.name='Times New Roman'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(4); style.paragraph_format.line_spacing=1.08
        if name=='Heading 1': style.paragraph_format.space_before=Pt(10)
        for el in list(style.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
        fonts=style.element.get_or_add_rPr().rFonts
        if fonts is not None:
            for attr in list(fonts.attrib):
                if 'Theme' in attr: del fonts.attrib[attr]
    doc.add_paragraph(title,'Title')
    for kind,value in sections:
        if kind=='h': md+=['## '+value,'']; doc.add_paragraph(value,'Heading 1')
        elif kind=='p': md+=[value,'']; doc.add_paragraph(value)
        else:
            md+=['| '+' | '.join(value[0])+' |','| '+' | '.join(['---']*len(value[0]))+' |']
            md+=['| '+' | '.join(row)+' |' for row in value[1:]]; md+=['']
            t=doc.add_table(rows=0,cols=len(value[0])); t.style='Table Grid'
            for ri,row in enumerate(value):
                cells=t.add_row().cells
                for cell,text in zip(cells,row):
                    cell.text=text
                    for p0 in cell.paragraphs:
                        p0.paragraph_format.space_after=Pt(4)
                        for run in p0.runs:
                            run.font.size=Pt(10)
                            run.font.name='Times New Roman'
                            run.bold=ri==0
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');cell._tc.get_or_add_tcPr().append(shade)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph('')
    (out/'REPORT_task02.md').write_text('\n'.join(md),encoding='utf-8')
    doc.save(out/'REPORT_task02.docx');print(out/'REPORT_task02.docx')

if __name__=='__main__':main()
