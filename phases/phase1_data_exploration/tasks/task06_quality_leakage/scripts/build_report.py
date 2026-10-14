"""Build Task 6 Markdown and Word reports from a completed run."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = next(parent for parent in Path(__file__).resolve().parents if (parent / "AGENTS.md").is_file())


def fmt(value, digits=0): return f"{value:,.{digits}f}"
def pct(value, total): return f"{100*value/total:.3f}%" if total else "—"
def md_table(headers, rows): return ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"] + ["| " + " | ".join(map(str, row)) + " |" for row in rows]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = tc_pr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_margins(cell, value=85):
    tc_pr = cell._tc.get_or_add_tcPr(); tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None: tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None: node = OxmlElement(f"w:{name}"); tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = str(header); cell.width = Cm(widths[i]); shade(cell, "1F4E78"); set_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(font_size); run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, row)):
            cell.text = str(value); cell.width = Cm(widths[i]); set_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2: shade(cell, "EAF2F8")
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (0, len(cells)-1) else WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.0
            for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(font_size)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el=OxmlElement(f"w:{edge}"); el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"4"); el.set(qn("w:color"),"D9D9D9"); borders.append(el)
    table._tbl.tblPr.append(borders); return table


def build_docx(path, sections):
    template = ROOT / "tmp" / "phase1_reference.docx"; doc = Document(template) if template.exists() else Document(); body=doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"): body.remove(child)
    for section in doc.sections:
        section.top_margin=Cm(2); section.bottom_margin=Cm(2); section.left_margin=Cm(2.2); section.right_margin=Cm(1.8)
        for part in (section.header,section.footer,section.first_page_header,section.first_page_footer):
            for p in part.paragraphs: p.clear()
    for name,size in (("Normal",11),("Title",18),("Heading 1",14),("Heading 2",12)):
        style=doc.styles[name]; style.font.name="Arial"; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0); style.paragraph_format.space_after=Pt(7); style.paragraph_format.line_spacing=1.2
    doc.styles["Normal"].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for entry in sections:
        kind=entry[0]
        if kind=="title": doc.add_paragraph(entry[1],"Title")
        elif kind=="subtitle": p=doc.add_paragraph(entry[1]); p.paragraph_format.space_after=Pt(12)
        elif kind=="h1": doc.add_paragraph(entry[1],"Heading 1")
        elif kind=="p": doc.add_paragraph(entry[1])
        elif kind=="bullet": p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.5); p.add_run("• "+entry[1])
        elif kind=="table": add_table(doc,entry[1],entry[2],entry[3],entry[4] if len(entry)>4 else 8.5)
    doc.core_properties.title="Báo cáo Task 6 chất lượng dữ liệu và leakage"; doc.core_properties.author=""; doc.save(path)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument("run_dir",type=Path); args=parser.parse_args(); run=args.run_dir.resolve()
    data=json.loads((run/"results.json").read_text(encoding="utf-8")); cohort=pd.read_csv(run/"cohort_readiness.csv"); quality=pd.read_csv(run/"quality_issue_register.csv"); leakage=pd.read_csv(run/"leakage_control_plan.csv")
    cohort_names={"all_admissions":"Toàn bộ admissions","diagnosis_available":"Có diagnosis","candidate_lab_available":"Có lab ứng viên và diagnosis","lab_during_admission":"Có lab trong admission và diagnosis","lab_within_6h":"Có lab 6 giờ và diagnosis","lab_within_12h":"Có lab 12 giờ và diagnosis","lab_within_24h":"Có lab 24 giờ và diagnosis"}
    cohort_rows=[[cohort_names.get(r.cohort_step,r.cohort_step),fmt(r.admissions),fmt(r.subjects),fmt(r.subjects_with_multiple_admissions),f"{100*r.share_all_admissions:.3f}%"] for r in cohort.itertuples()]
    quality_rows=[[r.source,r.issue,fmt(r.count),fmt(r.denominator),f"{100*r.rate:.3f}%" if pd.notna(r.rate) else "—",r.decision_or_next_action] for r in quality.itertuples()]
    leakage_rows=[[r.risk,r.mechanism,r.required_control,r.decision_status] for r in leakage.itertuples()]
    md=[
        "# Báo cáo Task 6 chất lượng dữ liệu và nguy cơ leakage","","## 1 Mục đích và phạm vi","",
        "Task 6 tổng hợp các phát hiện đã đo trong Task 2–5, mô tả số admission còn lại ở từng điều kiện và lập kế hoạch kiểm soát leakage cho các phase sau. Task không tạo dataset huấn luyện, không chọn cohort cuối, không chọn cửa sổ chính thức và không chọn phenotype.","",
        "## 2 Mức sẵn sàng của cohort","",*md_table(["Điều kiện","Admission","Subject","Subject có nhiều admission","Tỷ lệ toàn admission"],cohort_rows),"",
        "Các hàng 6, 12 và 24 giờ là ba kịch bản mô tả độc lập. Chúng chưa phải ba cohort được lựa chọn. Việc chia dữ liệu sau này phải thực hiện theo subject_id vì nhiều bệnh nhân có hơn một admission đủ điều kiện.","",
        "## 3 Danh mục vấn đề chất lượng","",*md_table(["Nguồn","Vấn đề","Số lượng","Mẫu số","Tỷ lệ","Quyết định hoặc bước tiếp"],quality_rows),"",
        "Dữ liệu gốc không bị sửa. Các quyết định loại dòng chỉ được áp dụng khi tạo dataset xử lý mới và phải kèm audit count.","",
        "## 4 Nguy cơ leakage và kiểm soát","",*md_table(["Nguy cơ","Cơ chế","Kiểm soát bắt buộc","Trạng thái"],leakage_rows),"",
        "## 5 Các quyết định đã chốt","","- Không khôi phục hadm_id còn thiếu.","- Tổng hợp nhánh feature và nhánh label riêng ở cấp hadm_id rồi mới ghép.","- Chia train, validation và test theo subject_id.","- Không diễn giải diagnoses_icd là thời điểm bác sĩ xác lập chẩn đoán.","",
        "## 6 Các quyết định còn chờ","","- Chọn cửa sổ giai đoạn sớm sau khi so sánh 6, 12 và 24 giờ.","- Quy tắc xử lý 10 itemid có nhiều đơn vị.","- Quy tắc dùng kết quả chữ và xử lý outlier.","- Mapping ICD sang Phecode và ngưỡng tần suất phenotype.","- Inclusion và exclusion criteria của cohort cuối.","",
        "## 7 Kết luận Task 6","",
        "Phase 1 đã xác định được các điểm có thể dùng làm nền tảng tạo cohort và các vấn đề phải kiểm soát. Chưa có căn cứ để tuyên bố dataset cuối đã sạch hoặc sẵn sàng train. Task 7 cần tổng hợp quyết định chuyển phase sau khi người nghiên cứu chọn các quy tắc còn mở.","",
        "## 8 File đầu ra","","- `cohort_readiness.csv`: số admission và subject theo từng điều kiện.","- `quality_issue_register.csv`: danh mục vấn đề chất lượng và cách xử lý.","- `leakage_control_plan.csv`: nguy cơ leakage và kiểm soát bắt buộc.","- `results.json`: metadata và số liệu máy đọc.","- `task06_quality_review.xlsx`: workbook dành cho người đọc.",
    ]
    (run/"REPORT_task06.md").write_text("\n".join(md),encoding="utf-8")
    sections=[
        ("title","BÁO CÁO TASK 6 CHẤT LƯỢNG DỮ LIỆU VÀ LEAKAGE"),("subtitle","Giai đoạn 1 nghiên cứu dự đoán đa nhãn từ dữ liệu xét nghiệm MIMIC IV"),
        ("p","Báo cáo tổng hợp chất lượng dữ liệu, mức sẵn sàng của cohort và các biện pháp kiểm soát leakage dựa trên kết quả Task 2 đến Task 5. Chưa có dataset huấn luyện hoặc cohort cuối được tạo."),
        ("h1","1 Mục đích và phạm vi"),("p","Task 6 đối chiếu các nhánh admission, xét nghiệm và chẩn đoán ở cấp hadm_id. Báo cáo giữ riêng sự kiện đã đo, quyết định đã chốt và lựa chọn nghiên cứu còn chờ."),
        ("h1","2 Mức sẵn sàng của cohort"),("table",["Điều kiện","Admission","Subject","Nhiều admission","Tỷ lệ"],cohort_rows,[6.0,2.5,2.5,3.0,2.5],8.2),
        ("p","Ba hàng cửa sổ là so sánh mô tả. Việc chia dữ liệu phải dùng subject_id để admission của cùng bệnh nhân không xuất hiện ở nhiều tập."),
        ("h1","3 Danh mục vấn đề chất lượng"),("table",["Nguồn","Vấn đề","Số lượng","Mẫu số","Tỷ lệ","Quyết định hoặc bước tiếp"],quality_rows,[2.2,4.1,2.4,2.4,2.0,5.0],7.5),
        ("p","Dữ liệu HOSP gốc không bị sửa. Những dòng không đủ điều kiện chỉ bị loại khi tạo dataset xử lý mới và phải có thống kê audit."),
        ("h1","4 Nguy cơ leakage và kiểm soát"),("table",["Nguy cơ","Cơ chế","Kiểm soát bắt buộc","Trạng thái"],leakage_rows,[3.3,5.0,6.0,3.2],7.8),
        ("h1","5 Các quyết định đã chốt"),("bullet","Không khôi phục hadm_id còn thiếu."),("bullet","Tổng hợp X và Y riêng ở cấp hadm_id rồi mới ghép."),("bullet","Chia train, validation và test theo subject_id."),("bullet","Không xem diagnoses_icd là timestamp chẩn đoán."),
        ("h1","6 Các quyết định còn chờ"),("bullet","Cửa sổ giai đoạn sớm."),("bullet","Chuẩn hóa đơn vị, kết quả chữ và outlier."),("bullet","Mapping ICD sang Phecode và ngưỡng phenotype."),("bullet","Cohort cuối và inclusion exclusion criteria."),
        ("h1","7 Kết luận"),("p","Task 6 đã lập danh mục vấn đề và kiểm soát cần thiết. Dataset chưa sẵn sàng train cho đến khi các quyết định cohort, cửa sổ, feature và label được hoàn tất ở các phase tiếp theo."),
        ("h1","8 File đầu ra"),("p","Kết quả gồm cohort_readiness.csv, quality_issue_register.csv, leakage_control_plan.csv, results.json và task06_quality_review.xlsx."),
    ]
    build_docx(run/"REPORT_task06.docx",sections); print(run/"REPORT_task06.docx")


if __name__=="__main__": main()
